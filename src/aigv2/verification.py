import io

from openpyxl import Workbook
from openpyxl.styles import Font

from src.aigv2.planning.meals import Exclusion
from src.aigv2.qc.facts import FactRow

_PLACES_HEADER = ["Day", "Place", "Guide says: hours", "Google says: hours", "Guide says: travel",
                  "Google says: travel", "Maps link", "Match", "Checked by"]
_EXCLUDED_HEADER = ["Name", "Reason", "Detail"]
_WIDTHS = {"Day": 6, "Place": 32, "Guide says: hours": 26, "Google says: hours": 30, "Guide says: travel": 22,
           "Google says: travel": 22, "Maps link": 48, "Match": 8, "Checked by": 18}


def _header(ws, columns: list[str]):
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = _WIDTHS.get(cell.value, 24)
    ws.freeze_panes = "A2"


def _places_sheet(ws, rows: list[FactRow]):
    _header(ws, _PLACES_HEADER)
    for r in rows:
        ws.append([r.day, r.venue, r.guide_hours, r.google_hours, r.guide_travel, r.google_travel,
                   r.link, "✓" if r.match else "✗", ""])
        if r.link:
            ws.cell(row=ws.max_row, column=7).hyperlink = r.link


def _excluded_sheet(ws, exclusions: list[Exclusion]):
    _header(ws, _EXCLUDED_HEADER)
    for e in exclusions:
        ws.append([e.name, e.reason, e.detail])


def build_sheet(rows: list[FactRow], exclusions: list[Exclusion]) -> bytes:
    wb = Workbook()
    _places_sheet(wb.active, rows)
    wb.active.title = "Places"
    _excluded_sheet(wb.create_sheet("Excluded restaurants"), exclusions)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
