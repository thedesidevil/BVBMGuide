from pathlib import Path

from docx import Document

from src.aigv2.intake.template import FIELDS, build_template, parse_template

REQUIRED = [label for label, req in FIELDS if req]


def _fill(path: Path, answers: dict[str, str]) -> None:
    doc = Document(str(path))
    for row in doc.tables[0].rows[1:]:
        label = row.cells[0].text.replace("*", "").strip()
        if label in answers:
            row.cells[1].text = answers[label]
    doc.save(str(path))


def test_build_then_parse_empty(tmp_path):
    p = tmp_path / "t.docx"
    build_template(p)
    result = parse_template(p)
    assert result.is_template is True
    assert result.missing_required == REQUIRED
    assert len(REQUIRED) == 6


def test_parse_filled(tmp_path):
    p = tmp_path / "t.docx"
    build_template(p)
    _fill(p, {
        "Client name(s)": "Saumitra",
        "Itinerary": "see attachments",
        "Diet (restrictions, allergies, likes)": "no beef, nothing raw",
    })
    result = parse_template(p)
    assert result.values["Client name(s)"] == "Saumitra"
    assert "Itinerary" in result.see_attachments
    assert "Itinerary" not in result.missing_required
    assert "Client name(s)" not in result.missing_required
    assert result.values["Diet (restrictions, allergies, likes)"] == "no beef, nothing raw"


def test_see_attachments_variants_case_insensitive(tmp_path):
    p = tmp_path / "t.docx"
    build_template(p)
    _fill(p, {"Itinerary": "  See Attached ", "Hotel per city": "USE INFORMATION FROM THE ATTACHED FILES"})
    result = parse_template(p)
    assert {"Itinerary", "Hotel per city"} <= result.see_attachments


def test_freeform_notes_not_template(tmp_path):
    p = tmp_path / "notes.docx"
    doc = Document()
    doc.add_heading("Portugal | 14–18 October", level=3)
    doc.add_paragraph("2 adults staying at Lisbon")
    doc.add_paragraph("Apart from pork they eat everything")
    doc.add_paragraph("Day 1: Arrive, check in, evening walk")
    doc.add_paragraph("Day 2: Sintra day trip")
    doc.save(str(p))
    result = parse_template(p)
    assert result.is_template is False
    assert result.values == {}
