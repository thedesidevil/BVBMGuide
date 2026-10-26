import os
from datetime import date, datetime
from pathlib import Path

from docx import Document

from src.library.ingest_core.ordering import (
    FileScan,
    OrderedFile,
    order_files,
    scan_file,
    write_order_report,
)


def _set_mtime(path: Path, d: date) -> None:
    ts = datetime(d.year, d.month, d.day, 12, 0, 0).timestamp()
    os.utime(path, (ts, ts))


def _new_doc(path: Path, lines: list[str], mtime: date | None = None) -> Path:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    doc.save(path)
    if mtime:
        _set_mtime(path, mtime)
    return path


def _old_doc(path: Path, rows: list[list[str]], mtime: date | None = None) -> Path:
    doc = Document()
    doc.add_paragraph("All Inclusive Guide")
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    doc.save(path)
    if mtime:
        _set_mtime(path, mtime)
    return path


def _scan(path, mtime, fmt="new", last_day=None, looks=True, h=None):
    return FileScan(
        path=Path(path),
        mtime=mtime,
        text_hash=h or f"hash-{Path(path).name}",
        fmt=fmt,
        last_day=last_day,
        looks_like_aig=looks,
    )


# (a) bulk mtime detection
def test_bulk_mtime_not_trusted_unique_mtimes_kept(tmp_path):
    bulk = date(2026, 3, 1)
    scans = []
    for i in range(7):
        p = _new_doc(tmp_path / f"bulk{i}.docx", [f"Day 1 trip {i}", f"Day 2 trip {i}"], bulk)
        scans.append(scan_file(p))
    u1 = _new_doc(tmp_path / "u1.docx", ["Day 1 one", "Day 2 one"], date(2025, 5, 1))
    u2 = _new_doc(tmp_path / "u2.docx", ["Day 1 two", "Day 2 two"], date(2025, 6, 1))
    scans += [scan_file(u1), scan_file(u2)]

    ordered = order_files(scans)
    by_name = {o.scan.path.name: o for o in ordered}
    for i in range(7):
        o = by_name[f"bulk{i}.docx"]
        assert o.basis != "mtime"
        assert o.order_date is None
    assert by_name["u1.docx"].basis == "mtime"
    assert by_name["u1.docx"].order_date == date(2025, 5, 1)
    assert by_name["u2.docx"].basis == "mtime"
    assert by_name["u2.docx"].order_date == date(2025, 6, 1)


def test_bulk_file_uses_travel_date_when_available(tmp_path):
    bulk = date(2026, 3, 1)
    scans = [_scan(f"b{i}.docx", bulk) for i in range(6)]
    scans.append(_scan("travel.docx", bulk, last_day=date(2024, 9, 3)))
    ordered = order_files(scans)
    o = next(o for o in ordered if o.scan.path.name == "travel.docx")
    assert o.basis == "travel"
    assert o.order_date == date(2024, 9, 3)


def test_mtime_shared_by_threshold_files_is_still_trusted():
    d = date(2025, 1, 1)
    scans = [_scan(f"f{i}.docx", d) for i in range(5)]
    ordered = order_files(scans, bulk_threshold=5)
    assert all(o.basis == "mtime" and o.order_date == d for o in ordered)


# (b) year inference from weekday headers
def test_last_day_from_weekday_headers(tmp_path):
    p = _new_doc(
        tmp_path / "w.docx",
        [
            "Day 1: Mon, 12 Jan – Arrival",
            "Day 2: Tue, 13 Jan – Old town",
            "Day 3: Wed, 14 Jan – Departure",
        ],
    )
    assert scan_file(p).last_day == date(2026, 1, 14)


def test_last_day_month_first_order(tmp_path):
    p = _new_doc(
        tmp_path / "w.docx",
        [
            "Day 1: Monday, Jan 12 – Arrival",
            "Day 2: Tuesday, Jan 13th – Old town",
        ],
    )
    assert scan_file(p).last_day == date(2026, 1, 13)


def test_last_day_long_form_with_ordinals(tmp_path):
    p = _new_doc(
        tmp_path / "w.docx",
        [
            "Day 1 – Monday 12th January",
            "Day 2 – Tuesday 13th January",
        ],
    )
    assert scan_file(p).last_day == date(2026, 1, 13)


def test_last_day_none_with_single_header(tmp_path):
    p = _new_doc(tmp_path / "w.docx", ["Day 1: Mon, 12 Jan", "Day 2: Arrival"])
    assert scan_file(p).last_day is None


def test_last_day_none_when_weekdays_inconsistent(tmp_path):
    p = _new_doc(
        tmp_path / "w.docx",
        ["Day 1: Mon, 12 Jan", "Day 2: Mon, 13 Jan", "Day 3: Mon, 14 Jan"],
    )
    assert scan_file(p).last_day is None


def test_feb_29_skipped_in_non_leap_years(tmp_path):
    # 2024-02-28 is Wednesday, 2024-02-29 Thursday, 2024-03-01 Friday
    p = _new_doc(
        tmp_path / "w.docx",
        ["Day 1: Wed, 28 Feb", "Day 2: Thu, 29 Feb", "Day 3: Fri, 1 Mar"],
    )
    assert scan_file(p).last_day == date(2024, 3, 1)


# (c) old-format detection
def test_three_row_day_table_is_old(tmp_path):
    p = _old_doc(
        tmp_path / "old.docx",
        [["Day 1", "Arrive"], ["Day 2", "Museum"], ["Day 3", "Leave"]],
    )
    s = scan_file(p)
    assert s.fmt == "old"
    assert s.looks_like_aig is True


def test_two_row_day_table_is_not_old(tmp_path):
    p = _old_doc(tmp_path / "x.docx", [["Day 1", "Arrive"], ["Day 2", "Leave"]])
    assert scan_file(p).fmt == "new"


def test_paragraph_days_are_new(tmp_path):
    p = _new_doc(tmp_path / "n.docx", ["Day 1: Arrival", "Day 2: Tour", "Day 3: Leave"])
    s = scan_file(p)
    assert s.fmt == "new"
    assert s.looks_like_aig is True


# (d) format outranks date
def test_old_format_2026_sorts_before_new_format_2024():
    scans = [
        _scan("new2024.docx", date(2024, 1, 1), fmt="new"),
        _scan("old2026.docx", date(2026, 6, 1), fmt="old"),
    ]
    ordered = order_files(scans)
    assert [o.scan.path.name for o in ordered] == ["old2026.docx", "new2024.docx"]


# (e) duplicates
def test_identical_text_second_is_duplicate(tmp_path):
    a = _new_doc(tmp_path / "a.docx", ["Day 1 same", "Day 2 same"], date(2025, 1, 1))
    b = _new_doc(tmp_path / "b.docx", ["Day 1 same", "Day 2 same"], date(2025, 2, 1))
    sa, sb = scan_file(a), scan_file(b)
    assert sa.text_hash == sb.text_hash
    ordered = order_files([sb, sa])
    by_name = {o.scan.path.name: o for o in ordered}
    assert by_name["a.docx"].duplicate_of is None
    assert by_name["a.docx"].skipped_reason is None
    assert by_name["b.docx"].duplicate_of == a
    assert by_name["b.docx"].skipped_reason == "duplicate"
    # skipped entries come after processable ones
    assert [o.scan.path.name for o in ordered] == ["a.docx", "b.docx"]


def test_emoji_prefixed_day_lines_look_like_aig(tmp_path):
    p = _new_doc(tmp_path / "e.docx", ["📅 Day 1: May 23 — Arrival", "🌏 Day 2: May 24 — Tour"])
    assert scan_file(p).looks_like_aig is True


# (f) not an AIG
def test_doc_without_day_lines_is_not_an_aig(tmp_path):
    p = _new_doc(tmp_path / "process.docx", ["Step 1: open the file", "Step 2: edit"])
    s = scan_file(p)
    assert s.looks_like_aig is False
    o = order_files([s])[0]
    assert o.skipped_reason == "not_an_aig"


def test_unreadable_file_does_not_raise(tmp_path):
    p = tmp_path / "bad.docx"
    p.write_bytes(b"not a docx")
    s = scan_file(p)
    assert s.looks_like_aig is False
    assert s.fmt == "new"
    assert s.text_hash
    assert s.last_day is None


def test_skipped_sort_after_processable():
    scans = [
        _scan("a.docx", date(2025, 1, 1), looks=False),
        _scan("z.docx", date(2025, 2, 1)),
    ]
    ordered = order_files(scans)
    assert [o.scan.path.name for o in ordered] == ["z.docx", "a.docx"]


# (g) determinism
def test_order_is_deterministic():
    bulk = date(2026, 3, 1)
    scans = [_scan(f"f{i}.docx", bulk, fmt="old" if i % 2 else "new") for i in range(8)]
    scans.append(_scan("m.docx", date(2025, 1, 1)))
    first = [o.scan.path for o in order_files(scans)]
    second = [o.scan.path for o in order_files(list(reversed(scans)))]
    assert first == second


def test_basis_format_and_name():
    bulk = date(2026, 3, 1)
    scans = [_scan(f"n{i}.docx", bulk, fmt="new") for i in range(6)]
    scans.append(_scan("old.docx", bulk, fmt="old"))
    ordered = order_files(scans)
    by_name = {o.scan.path.name: o for o in ordered}
    assert by_name["old.docx"].basis == "format"
    assert all(by_name[f"n{i}.docx"].basis == "name" for i in range(6))
    assert [o.scan.path.name for o in ordered] == ["old.docx"] + [f"n{i}.docx" for i in range(6)]


def test_write_order_report(tmp_path):
    bulk = date(2026, 3, 1)
    scans = [_scan(f"n{i}.docx", bulk) for i in range(6)]
    scans.append(_scan("dated.docx", date(2025, 1, 1)))
    scans.append(_scan("junk.docx", date(2025, 2, 2), looks=False))
    ordered = order_files(scans)
    out = tmp_path / "order.md"
    write_order_report(ordered, out)
    text = out.read_text()
    assert "| position | file | order_date | basis | format | skipped reason |" in text
    assert "dated.docx" in text and "2025-01-01" in text
    assert "not_an_aig" in text
    closing = text.split("## ", 2)[-1]
    for i in range(6):
        assert f"n{i}.docx" in closing
    assert "dated.docx" not in closing
    assert "junk.docx" not in closing
