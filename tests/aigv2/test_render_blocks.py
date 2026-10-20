from docx.shared import Cm
from docx.oxml.ns import qn

from src.aigv2.models import Marker
from src.aigv2.render import blocks
from src.aigv2.render.theme import new_document


def test_new_document_is_a4():
    s = new_document().sections[0]
    assert abs(s.page_width.cm - 21.0) < 0.01
    assert abs(s.page_height.cm - 29.7) < 0.01
    assert abs(s.left_margin.cm - 2.0) < 0.01


def test_heading_styles_builtin_and_branded():
    st = new_document().styles
    for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        assert st[name].font.name == "Cormorant Garamond"
    assert str(st["Heading 1"].font.color.rgb) == "0C1F3E"
    assert str(st["Heading 2"].font.color.rgb) == "C68A32"
    assert st["Heading 2"].font.italic is True
    assert st["Normal"].font.name == "Georgia"


def test_link_is_real_hyperlink():
    doc = new_document()
    p = doc.add_paragraph()
    blocks.add_link(p, "Map", "https://example.com/x")
    h = p._p.find(qn("w:hyperlink"))
    assert h is not None
    rid = h.get(qn("r:id"))
    assert doc.part.rels[rid].target_ref == "https://example.com/x"
    assert doc.part.rels[rid].is_external


def test_marker_highlighted_yellow():
    doc = new_document()
    p = doc.add_paragraph()
    m = Marker(kind="VERIFY", text="hours")
    blocks.add_marker(p, m)
    assert p.runs[0].text == m.render()
    assert p._p.xpath('.//w:highlight[@w:val="yellow"]')


def test_card_structure():
    doc = new_document()
    blocks.add_card(
        doc, "Little Venice", "https://example.com",
        [("Opening hours", "Open 24 hours"), ("Entry", Marker(kind="PRICE_UNKNOWN", text="")),
         ("Must try", ["a", "b"])],
        "A quiet corner.", ["Vegetarian", "Casual lunch"],
    )
    assert len(doc.tables) == 1
    title = doc.tables[0].cell(0, 0).paragraphs[0]
    assert title.style.name == "Heading 3"
    assert title._p.find(qn("w:hyperlink")) is not None
    text = "".join(t.text for t in doc.tables[0]._tbl.iter(qn("w:t")))
    for s in ("Opening hours", "Open 24 hours", "⚠ PRICE UNKNOWN", "Vegetarian", "A quiet corner."):
        assert s in text
    assert doc.tables[0]._tbl.xpath('.//w:shd[@w:fill="FBFDFF"]')
    assert doc.tables[0]._tbl.xpath('.//w:shd[@w:fill="EEF4FB"]')


def test_card_without_url_and_facts():
    doc = new_document()
    blocks.add_card(doc, "Plain", None, [], None, [])
    assert doc.tables[0].cell(0, 0).paragraphs[0].style.name == "Heading 3"


def test_slot_heading_eyebrow_rule_bookmark():
    doc = new_document()
    assert blocks.add_slot_heading(doc, "Morning").style.name == "Heading 2"
    blocks.add_eyebrow(doc, "Day 2")
    blocks.add_rule(doc)
    p = doc.add_paragraph()
    blocks.add_bookmark(p, "day_2")
    q = doc.add_paragraph()
    blocks.add_internal_link(q, "Day 2", "day_2")
    assert p._p.xpath('.//w:bookmarkStart[@w:name="day_2"]')
    assert q._p.xpath('.//w:hyperlink[@w:anchor="day_2"]')


def test_footer_has_tagline_and_page_field():
    doc = new_document()
    f = doc.sections[0].footer._element
    text = "".join(t.text for t in f.iter(qn("w:t")))
    assert "Your Vacation · Your Story · Your Way" in text
    instr = "".join(t.text for t in f.iter(qn("w:instrText")))
    assert "PAGE" in instr


def test_header_has_logo():
    doc = new_document()
    h = doc.sections[0].header._element
    assert h.xpath(".//pic:pic") or h.xpath(".//*[local-name()='pic']")


from src.aigv2.render.theme import PPR_ORDER, TBLPR_ORDER, TCPR_ORDER


def _local(el):
    return [c.tag.split("}")[1] for c in el]


def _assert_ordered(el, order):
    idx = [order.index(n) for n in _local(el)]
    assert idx == sorted(idx), _local(el)


def test_card_xml_child_order():
    doc = new_document()
    blocks.add_card(doc, "T", "https://e.com", [("A", "b")], "body", ["p"])
    tbl = doc.tables[0]._tbl
    _assert_ordered(tbl.tblPr, TBLPR_ORDER)
    for tcpr in tbl.iter(qn("w:tcPr")):
        _assert_ordered(tcpr, TCPR_ORDER)
    for ppr in tbl.iter(qn("w:pPr")):
        _assert_ordered(ppr, PPR_ORDER)


def test_rule_and_footer_ppr_order():
    doc = new_document()
    p = blocks.add_rule(doc)
    _assert_ordered(p._p.pPr, PPR_ORDER)
    assert "pBdr" in _local(p._p.pPr)
    _assert_ordered(doc.sections[0].footer.paragraphs[0]._p.pPr, PPR_ORDER)


def test_links_not_underlined_and_no_blank_paragraph_in_card():
    doc = new_document()
    p = doc.add_paragraph()
    blocks.add_link(p, "x", "https://e.com")
    assert p._p.xpath('.//w:u[@w:val="none"]')
    blocks.add_card(doc, "T", "https://e.com", [("A", "b")], "body", [])
    cell = doc.tables[0].cell(0, 0)
    assert [q.text for q in cell.paragraphs] == ["T", "body"]
