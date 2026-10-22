from typing import Iterable, Optional
from urllib.parse import quote_plus

from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.aigv2.models import Marker, Place
from src.aigv2.planning.important_places import ImportantPlaces, NearbyPlace
from src.aigv2.render.blocks import add_bookmark, add_card, add_eyebrow, add_link, add_marker
from src.aigv2.render.theme import ASSETS
from src.aigv2.writing.sections import SectionContent, SectionItem

WEBSITE = "https://www.bonvoyagebymarina.com"
_WEBSITE_TEXT = "www.bonvoyagebymarina.com"


def heading1(doc, text: str, bookmark: str):
    p = doc.add_paragraph(text, style="Heading 1")
    p.paragraph_format.page_break_before = True
    add_bookmark(p, bookmark)
    return p


def marker_lines(doc, markers: Iterable[Marker]) -> None:
    for m in markers:
        add_marker(doc.add_paragraph(), m)


def place_markers(place: Optional[Place], existing: Iterable[Marker] = ()) -> list[Marker]:
    if place is None or not place.verify_note:
        return []
    if any(m.text == place.verify_note for m in existing):
        return []
    return [Marker(kind="VERIFY", text=place.verify_note)]


def search_url(name: str, city: str = "") -> str:
    return f"https://www.google.com/maps/search/?api=1&query={quote_plus(f'{name} {city}'.strip())}"


def minutes_text(m: int) -> str:
    h, mi = divmod(int(m), 60)
    if not h:
        return f"{mi} min"
    return f"{h} h" + (f" {mi} min" if mi else "")


def _items(doc, items: list[SectionItem]) -> None:
    for it in items:
        doc.add_paragraph(it.heading, style="Heading 3")
        doc.add_paragraph(it.body)
        if it.verify:
            marker_lines(doc, [Marker(kind="VERIFY", text=it.heading)])


def render_section(doc, title: str, bookmark: str, content: SectionContent) -> None:
    heading1(doc, title, bookmark)
    if content.by_country:
        for country, items in content.by_country.items():
            doc.add_paragraph(country or "General", style="Heading 2")
            _items(doc, items)
    else:
        _items(doc, content.items)
    marker_lines(doc, content.markers)


def _nearby_card(doc, n: NearbyPlace, city: str) -> None:
    p = n.place
    travel = [f"Walk · {n.walk_min} min"] if n.walk_min is not None else []
    if n.drive_min is not None:
        travel.append(f"By car · {n.drive_min} min")
    facts = [("GETTING THERE", travel or Marker(kind="VERIFY", text="travel time from the hotel")),
             ("OPENING HOURS", n.hours_text)]
    if p.phone:
        facts.append(("PHONE", p.phone))
    if p.address:
        facts.append(("ADDRESS", p.address))
    add_card(doc, p.name, p.maps_url or search_url(p.name, city), facts, None, [])
    marker_lines(doc, n.markers + place_markers(p, n.markers))


def render_important(doc, title: str, bookmark: str, entries: list[ImportantPlaces]) -> None:
    heading1(doc, title, bookmark)
    for e in entries:
        doc.add_paragraph(f"{e.hotel}, {e.city}", style="Heading 2")
        for label, group in (("PHARMACIES", e.pharmacies), ("GROCERIES", e.groceries),
                             ("HOSPITAL", [e.hospital] if e.hospital else [])):
            if not group:
                continue
            add_eyebrow(doc, label)
            for n in group:
                _nearby_card(doc, n, e.city)
        marker_lines(doc, e.markers)


def render_thank_you(doc, destination: str) -> None:
    heading1(doc, "Thank You", "thank_you")
    text = (ASSETS / "thank_you.txt").read_text(encoding="utf-8").replace("{destination}", destination)
    for line in text.splitlines():
        if not line.strip():
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if _WEBSITE_TEXT in line:
            before, after = line.split(_WEBSITE_TEXT, 1)
            p.add_run(before)
            add_link(p, _WEBSITE_TEXT, WEBSITE)
            p.add_run(after)
        else:
            p.add_run(line)
