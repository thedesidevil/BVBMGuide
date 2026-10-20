import io
import re
from datetime import date, datetime, time

from docx import Document
from docx.oxml.ns import qn

from src.aigv2.config import Settings
from src.aigv2.models import (
    Booking, DayPlan, Diet, HotelStay, Marker, PlannedItem, Place, Stop, TransitStep, TransportLeg, Traveller,
    TripPlan, WeeklyHours,
)
from src.aigv2.planning.important_places import ImportantPlaces, NearbyPlace
from src.aigv2.planning.meals import ChosenRestaurant, MealPick
from src.aigv2.planning.slots import DayResult, MealSlot, ScheduledItem
from src.aigv2.planning.travel import Leg
from src.aigv2.render.document import (
    GuideData, RenderDay, day_header, output_filename, render_guide, sunset_time,
)
from src.aigv2.writing.days import AttractionText, DayText, RestaurantText
from src.aigv2.writing.sections import SectionContent, SectionItem

S = Settings(models={})
D1, D2, D3 = date(2026, 10, 29), date(2026, 10, 30), date(2026, 10, 31)
HOTEL = "Pera Palace"


def place(name, lat=41.03, lng=28.97, match_ok=True, website=None, note=None, hrs=True):
    hours = WeeklyHours(periods={d: [(time(9), time(18))] for d in range(7)}, source="google") if hrs else None
    return Place(query=name, place_id=f"pid_{name.replace(' ', '_')}", name=name, address="1 Test St",
                 lat=lat, lng=lng, status="OPERATIONAL", hours=hours, website=website, phone="+90 212 000 0000",
                 maps_url=f"https://www.google.com/maps/search/?api=1&query={name.replace(' ', '+')}"
                          f"&query_place_id=pid", fetched_at=datetime(2026, 10, 1), match_ok=match_ok,
                 verify_note=note)


def item(name, t=None, slot="any", kind="attraction", **kw):
    return PlannedItem(name=name, kind=kind, city="Istanbul", time=t, slot_hint=slot, **kw)


def sched(it, slot, start):
    return ScheduledItem(item=it, slot=slot, start=start)


def chosen(name, minutes=6, verify=False, **lib):
    return ChosenRestaurant(lib={"name": name, "cuisine_type": ["Turkish"], **lib}, place=place(name),
                            travel_min=minutes, travel_label="Walk", hours_text="9:00 AM – 6:00 PM",
                            hours_verify=verify)


FLIGHT_IN = TransportLeg(kind="flight", from_place="Mumbai", to_place="Istanbul Airport",
                         depart=datetime(2026, 10, 29, 4, 0), arrive=datetime(2026, 10, 29, 10, 30),
                         source="ticket", operator="Turkish Airlines", number="TK 721")
FLIGHT_OUT = TransportLeg(kind="flight", from_place="Istanbul Airport", to_place="Mumbai",
                          depart=datetime(2026, 10, 31, 14, 0), arrive=datetime(2026, 10, 31, 22, 0),
                          source="ticket", operator="Turkish Airlines", number="TK 720")


def trip(**kw):
    base = dict(
        client_names=["Annie Sharma", "Raj Sharma"], travellers=Traveller(adults=2), diet=Diet(),
        travel_modes=["taxi"],
        hotels=[HotelStay(city="Istanbul", hotel=HOTEL, check_in=D1, check_out=D3, source="voucher",
                          check_in_time=time(15))],
        transport=[FLIGHT_IN, FLIGHT_OUT], bookings=[], days=[
            DayPlan(date=D1, title="Welcome", city="Istanbul"), DayPlan(date=D2, title="Old City", city="Istanbul"),
            DayPlan(date=D3, title="Farewell", city="Istanbul")],
        open_choices=[], vague_days=[], commentary_removed=[], trip_scope="international")
    base.update(kw)
    return TripPlan(**base)


def day1():
    galata = item("Galata Tower", slot="afternoon")
    dinner = MealSlot(meal="dinner", start=time(19, 30), anchors=[(HOTEL, (41.0, 28.9), 3)])
    pick = MealPick(slot=dinner, day_index=1, restaurants=[chosen("Mikla"), chosen("Saray Muhallebicisi", verify=True)],
                    markers=[Marker(kind="ADD_RESTAURANT", text=f"Dinner, Day 1, near {HOTEL}")])
    return RenderDay(
        index=1, date=D1, city="Istanbul",
        text=DayText(title="Welcome to Istanbul", intro="We begin gently.",
                     attractions=[AttractionText(name="Galata Tower", description="A medieval stone tower.")],
                     restaurants=[RestaurantText(name="Mikla", ambience="Rooftop views", best_for="Special dinner")]),
        result=DayResult(scheduled=[sched(galata, "afternoon", time(16))], meals=[dinner]),
        meals=[pick],
        legs={"arrival": Leg(from_name="Istanbul Airport", to_name=HOTEL, drive_min=50),
              "Galata Tower": Leg(from_name=HOTEL, to_name="Galata Tower", walk_min=8, drive_min=5)},
        places={HOTEL: place(HOTEL), "Galata Tower": place("Galata Tower")},
        booking_by_item={}, price_by_restaurant={"Mikla": "€60 (~₹5,800)"},
        arrival=FLIGHT_IN, departure=None, prev_city=None)


def day2():
    topkapi, hagia = item("Topkapi Palace", t="10:00"), item("Hagia Sophia", slot="morning")
    cruise = item("Bosphorus Sunset Cruise", t="18:00", optional=True, note="Watch the sunset from the deck")
    lunch = MealSlot(meal="lunch", start=time(12, 30), included=True)
    dinner = MealSlot(meal="dinner", start=time(19, 30),
                      anchors=[("Bosphorus Sunset Cruise", (41.0, 29.0), 1), (HOTEL, (41.0, 28.9), 2)])
    pick = MealPick(slot=dinner, day_index=2,
                    restaurants=[chosen("Balikci Sabahattin"), chosen("Ciya Sofrasi"), chosen("Hamdi")])
    return RenderDay(
        index=2, date=D2, city="Istanbul",
        text=DayText(title="Sultans and Sunsets", intro="Today is about the old city.",
                     attractions=[AttractionText(name="Topkapi Palace", description="Ottoman palace.",
                                                 entry_fee="TRY 1500", dress_code="Shoulders covered"),
                                  AttractionText(name="Hagia Sophia", description="A great mosque."),
                                  AttractionText(name="Bosphorus Sunset Cruise", description="A gentle cruise.")]),
        result=DayResult(scheduled=[sched(topkapi, "morning", time(10)), sched(hagia, "morning", time(11, 30)),
                                    sched(cruise, "evening", time(18))],
                         meals=[lunch, dinner], tips=["Carry a scarf for mosque visits"],
                         markers=[Marker(kind="VERIFY", text="location of Spice Bazaar unknown")]),
        meals=[pick],
        legs={"Topkapi Palace": Leg(from_name=HOTEL, to_name="Topkapi Palace", walk_min=25, drive_min=15,
                                    transit_min=20, transit_steps=[
                                        TransitStep(line="T1", headsign="Kabatas", board="Karakoy", alight="Sultanahmet",
                                                    stops=4)])},
        places={HOTEL: place(HOTEL), "Topkapi Palace": place("Topkapi Palace", website="https://topkapi.example"),
                "Hagia Sophia": None, "Bosphorus Sunset Cruise": place("Bosphorus Sunset Cruise", note="matched by name only")},
        booking_by_item={"Topkapi Palace": Booking(name="Topkapi Palace", day=D2, time="10:00", source="voucher",
                                                   meeting_point="Main gate", instructions=["Show the QR code"]),
                         "Bosphorus Sunset Cruise": Booking(name="Bosphorus Sunset Cruise", day=D2, source="voucher",
                                                            included_in_package=True)},
        price_by_restaurant={}, arrival=None, departure=None, prev_city="Istanbul", utc_offset_hours=3.0)


def day3():
    breakfast = MealSlot(meal="breakfast", start=time(7, 30), included=True)
    return RenderDay(
        index=3, date=D3, city="Istanbul", text=DayText(title="Farewell, Istanbul", intro="Time to fly home."),
        result=DayResult(meals=[breakfast]), meals=[],
        legs={"departure": Leg(from_name=HOTEL, to_name="Istanbul Airport", drive_min=50)},
        places={HOTEL: place(HOTEL)}, booking_by_item={}, price_by_restaurant={},
        arrival=None, departure=FLIGHT_OUT, prev_city="Istanbul")


def sections():
    def sc(key, *headings, verify=False):
        return SectionContent(key=key, items=[SectionItem(heading=h, body=f"About {h}.", verify=verify) for h in headings])
    out = {k: sc(k, k.title()) for k in ("packing", "money", "connectivity", "health", "getting_around", "dishes",
                                         "souvenirs", "etiquette")}
    out["money"] = sc("money", "Cards", verify=True)
    out["safety"] = SectionContent(key="safety", items=[], by_country={
        "Turkey": [SectionItem(heading="Emergency 112", body="Call 112.")],
        "Greece": [SectionItem(heading="Emergency 112", body="Call 112 in Greece.")]})
    return out


def important():
    pharm = NearbyPlace(place=place("Eczane Pera"), walk_min=4, hours_text="9–6 Mon–Sun",
                        markers=[Marker(kind="VERIFY", text="opening hours unknown")])
    hosp = NearbyPlace(place=place("Amerikan Hastanesi"), drive_min=12, hours_text="24 hours Mon–Sun")
    return [ImportantPlaces(city="Istanbul", hotel=HOTEL, pharmacies=[pharm], hospital=hosp,
                            markers=[Marker(kind="VERIFY", text="fewer than 2 groceries within a 10-minute walk")])]


def data(**kw):
    base = dict(trip=trip(), title="Minarets and Moonlight", days=[day1(), day2(), day3()], sections=sections(),
                important=important(), trip_scope="international", modes=["taxi"], domestic=False)
    base.update(kw)
    return GuideData(**base)


def load(d=None):
    return Document(io.BytesIO(render_guide(d or data(), S)))


def all_text(doc):
    return "\n".join("".join(t.text or "" for t in p.iter(qn("w:t"))) for p in doc.element.body.iter(qn("w:p")))


def h1s(doc):
    return [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]


def day_block(doc, n):
    text = all_text(doc)
    start = text.index(f"Day {n} –", text.index(f"Day {n} –") + 1)
    nxt = text.find(f"Day {n + 1} –", start)
    return text[start:nxt if nxt > 0 else text.index("Tailored Packing List", start)]


def test_day_header_format():
    assert day_header(1, date(2026, 10, 29), "Welcome to Istanbul") == "Day 1 – Thu, 29 Oct | Welcome to Istanbul"
    assert day_header(5, date(2026, 11, 15), "X").startswith("Day 5 – Sun, 15 Nov | ")


def test_section_order():
    assert h1s(load()) == [
        "Day 1 – Thu, 29 Oct | Welcome to Istanbul", "Day 2 – Fri, 30 Oct | Sultans and Sunsets",
        "Day 3 – Sat, 31 Oct | Farewell, Istanbul",
        "Tailored Packing List", "Money & Payments", "Mobile Connectivity & SIM Guide", "Health & Vaccination",
        "Important Places Around Your Stays", "Getting Around", "Must-Try Local Dishes", "Souvenir Shopping Guide",
        "Cultural Etiquette & Local Phrases", "Safety & Emergency Contacts", "Thank You"]


def test_each_heading1_starts_new_page():
    doc = load()
    assert all(p.paragraph_format.page_break_before for p in doc.paragraphs if p.style.name == "Heading 1")


def test_toc_has_no_page_numbers_and_links_to_bookmarks():
    doc = load()
    body = doc.element.body
    marks = {}
    for p in doc.paragraphs:
        for b in p._p.findall(qn("w:bookmarkStart")):
            marks[b.get(qn("w:name"))] = p
    links = [h for h in body.iter(qn("w:hyperlink")) if h.get(qn("w:anchor"))]
    assert len(links) == len(h1s(doc)) - 1
    for h in links:
        target = marks[h.get(qn("w:anchor"))]
        assert target.style.name == "Heading 1"
        line = h.getparent()
        text = "".join(t.text for t in line.iter(qn("w:t")))
        assert not re.search(r"\d+\s*$", text) or text.startswith("Day")
        assert not line.findall(qn("w:r") + "/" + qn("w:tab"))
    assert "PAGEREF" not in doc.element.xml
    text = all_text(doc)
    for group in ("DAY BY DAY", "BEFORE YOU GO", "WHILE YOU'RE THERE"):
        assert group in text.upper()
    assert text.index("Your All-Inclusive Guide") < text.index("Minarets and Moonlight") < text.index(
        "All Inclusive Guide – Istanbul")


def test_letter_page_marker():
    text = all_text(load())
    assert "Dear Annie," in text
    assert "⚠ WRITE PERSONAL LETTER" in text
    assert text.index("Dear Annie,") < text.index("Minarets and Moonlight")


def test_cover_section_zero_margins_with_full_page_image():
    doc = load()
    cover = doc.sections[0]
    assert cover.left_margin == 0 and cover.top_margin == 0
    assert cover.different_first_page_header_footer
    assert abs(doc.sections[1].left_margin.cm - 2.0) < 0.01
    first = doc.paragraphs[0]
    assert first._p.xpath(".//pic:pic")


def test_every_venue_card_has_hyperlink():
    doc = load()
    cards = [t for t in doc.tables if t.cell(0, 0).paragraphs[0].style.name == "Heading 3"]
    assert len(cards) > 10
    for t in cards:
        assert t.cell(0, 0).paragraphs[0]._p.find(qn("w:hyperlink")) is not None, t.cell(0, 0).paragraphs[0].text


def test_no_utm_and_no_emoji():
    doc = load()
    text = all_text(doc)
    assert "utm_" not in doc.element.xml
    assert all("utm_" not in r.target_ref for r in doc.part.rels.values() if r.is_external)
    bad = [c for c in text if c != "⚠" and (0x1F000 <= ord(c) <= 0x1FAFF or 0x2600 <= ord(c) <= 0x27BF
                                           or ord(c) == 0xFE0F)]
    assert not bad


def test_output_filename():
    assert output_filename(trip()) == "All Inclusive Guide - Annie Sharma - Istanbul.docx"
    two = trip(hotels=[HotelStay(city="Paris", hotel="A", check_in=D1, check_out=D2, source="voucher"),
                       HotelStay(city="Rome", hotel="B", check_in=D2, check_out=D3, source="voucher"),
                       HotelStay(city="Paris", hotel="C", check_in=D3, check_out=D3, source="voucher")])
    assert output_filename(two) == "All Inclusive Guide - Annie Sharma - Paris, Rome.docx"


def test_untimed_booking_not_shown_as_confirmed_time():
    d2 = day2()
    d2.booking_by_item["Topkapi Palace"] = Booking(name="Topkapi Palace", day=D2, time="10:00", source="voucher",
                                                   timed=False)
    block = day_block(load(data(days=[day1(), d2, day3()])), 2)
    assert "YOUR BOOKING\nValid all day" in block
    assert "Confirmed" not in block
    timed = day_block(load(), 2)
    assert "YOUR BOOKING\n10:00" in timed
    assert "MEETING POINT\nMain gate" in timed
    assert "Show the QR code" in timed


def test_arrival_day_transfer_card_and_flight():
    block = day_block(load(), 1)
    assert "Istanbul Airport → Pera Palace" in block
    assert "By car · 50 min" in block
    assert "⚠ VERIFY: transfer cost" in block
    assert "TK 721" in block
    assert block.index("Istanbul Airport → Pera Palace") < block.index("Galata Tower")
    assert "Overnight in Istanbul" in block


def test_travel_line_before_card_and_transit_steps():
    block = day_block(load(data(modes=["taxi", "public_transport"])), 2)
    assert "Getting there: By car · 15 min  |  By public transport · 20 min" in block
    assert "T1, alight Sultanahmet, 4 stops" in block
    assert block.index("Getting there") < block.index("Topkapi Palace\n")


def test_attraction_ai_fields_marked_and_entry_fee_links_website():
    doc = load()
    block = day_block(doc, 2)
    assert "TRY 1500 · ⚠ VERIFY: entry fee" in block
    assert "Shoulders covered · ⚠ VERIFY: dress code" in block
    assert any(r.target_ref == "https://topkapi.example" for r in doc.part.rels.values() if r.is_external)


def test_missing_place_renders_search_link_and_verify():
    doc = load()
    block = day_block(doc, 2)
    seg = block[block.index("Hagia Sophia\n"):block.index("Bosphorus Sunset Cruise\n")]
    assert "⚠ VERIFY: place not found on Google" in seg
    assert "OPENING HOURS" not in seg
    hagia = next(t for t in doc.tables if t.cell(0, 0).paragraphs[0].text == "Hagia Sophia")
    rid = hagia.cell(0, 0).paragraphs[0]._p.find(qn("w:hyperlink")).get(qn("r:id"))
    url = doc.part.rels[rid].target_ref
    assert "query=Hagia+Sophia+Istanbul" in url and "query_place_id" not in url


def test_pills_tips_markers_and_sunset():
    block = day_block(load(), 2)
    seg = block[block.index("Bosphorus Sunset Cruise\n"):]
    assert "Optional" in seg and "Included in your package" in seg
    assert "⚠ VERIFY: matched by name only" in seg
    assert re.search(r"SUNSET\nabout 1[78]:\d\d", seg)
    assert "Carry a scarf for mosque visits" in block
    assert "⚠ VERIFY: location of Spice Bazaar unknown" in block


def test_meals_per_anchor_included_line_and_gap_markers():
    b2 = day_block(load(), 2)
    assert "Lunch — Included in your package" in b2
    assert "Lunch near" not in b2
    near_cruise = b2[b2.index("Dinner near Bosphorus Sunset Cruise"):b2.index(f"Dinner near {HOTEL}")]
    assert "Balikci Sabahattin" in near_cruise and "Ciya Sofrasi" not in near_cruise
    assert "Ciya Sofrasi" in b2[b2.index(f"Dinner near {HOTEL}"):]
    b1 = day_block(load(), 1)
    dinner = b1[b1.index(f"Dinner near {HOTEL}"):]
    assert f"⚠ ADD RESTAURANT – Dinner, Day 1, near {HOTEL}" in dinner
    assert "PRICE\n€60 (~₹5,800)" in dinner
    assert "⚠ PRICE UNKNOWN" in dinner
    assert "⚠ VERIFY: opening hours" in dinner
    assert "Rooftop views" in dinner
    assert "Walk · 6 min" in dinner


def test_departure_day_leave_by_and_no_overnight():
    block = day_block(load(), 3)
    assert "Breakfast — Included in your package" in block
    assert "Leave for the airport by 11:00" in block
    assert "⚠ VERIFY: departure time" in block
    assert block.index("Leave for the airport by") > block.index("Breakfast — Included")
    assert "Pera Palace → Istanbul Airport" in block
    assert "Overnight in" not in block


def test_domestic_flight_leave_by_two_hours():
    block = day_block(load(data(domestic=True, trip_scope="domestic")), 3)
    assert "Leave for the airport by 12:00" in block


def test_intercity_train_day():
    train = TransportLeg(kind="train", from_place="Paris Gare de Lyon", to_place="Lyon Part-Dieu",
                         depart=datetime(2026, 10, 30, 9, 0), arrive=datetime(2026, 10, 30, 11, 0),
                         source="ticket", ref="ABC123", operator="SNCF", number="TGV 6601", seat="Coach 5, Seat 42",
                         platform="7")
    t = trip(hotels=[HotelStay(city="Paris", hotel="Hotel P", check_in=D1, check_out=D2, source="voucher"),
                     HotelStay(city="Lyon", hotel="Hotel L", check_in=D2, check_out=D3, source="voucher",
                               check_in_time=time(15))])
    d = RenderDay(index=2, date=D2, city="Lyon", text=DayText(title="On to Lyon", intro="By rail today."),
                  result=DayResult(), meals=[], places={"Hotel P": place("Hotel P"), "Hotel L": place("Hotel L")},
                  legs={"departure": Leg(from_name="Hotel P", to_name="Paris Gare de Lyon", drive_min=20),
                        "arrival": Leg(from_name="Lyon Part-Dieu", to_name="Hotel L", walk_min=9)},
                  booking_by_item={}, price_by_restaurant={}, arrival=train, departure=train, prev_city="Paris")
    block = day_block(load(data(trip=t, days=[day1(), d])), 2)
    for s in ("Hotel P → Paris Gare de Lyon", "SNCF", "TGV 6601", "Coach 5, Seat 42", "09:00", "11:00", "2 h",
              "Reach the station by 08:20", "PLATFORM\n7", "Lyon Part-Dieu → Hotel L", "Walk · 9 min",
              "⚠ VERIFY: inter-city travel day logistics", "luggage", "Overnight in Lyon"):
        assert s in block, s
    assert block.index("Hotel P → Paris Gare de Lyon") < block.index("TGV 6601") < block.index("Lyon Part-Dieu → Hotel L")
    assert block.count("⚠ VERIFY: inter-city travel day logistics") == 1


def test_untickted_train_hides_platform_and_boat_arrival_label():
    train = TransportLeg(kind="train", from_place="A", to_place="B", depart=datetime(2026, 10, 30, 9, 0),
                         arrive=datetime(2026, 10, 30, 10, 0), source="notes", platform="3")
    boat = TransportLeg(kind="boat", from_place="Piraeus", to_place="", arrive=datetime(2026, 10, 29, 10, 0),
                        source="ticket")
    d1 = day1().model_copy(update={"arrival": boat})
    d2 = RenderDay(index=2, date=D2, city="Istanbul", text=DayText(title="T", intro="I"), result=DayResult(),
                   meals=[], legs={}, places={}, booking_by_item={}, price_by_restaurant={}, arrival=train,
                   departure=train, prev_city="Istanbul")
    doc = load(data(days=[d1, d2]))
    assert "PLATFORM" not in day_block(doc, 2)
    assert "Harbor → Pera Palace" in day_block(doc, 1)


def test_self_drive_parking_fact():
    block = day_block(load(data(modes=["self_drive"])), 1)
    galata = block[block.index("Galata Tower\n"):block.index("Dinner near")]
    assert "PARKING\n⚠ VERIFY: parking" in galata
    d1 = day1()
    d1.meals[0].restaurants[0].lib["parking"] = "Valet available"
    block = day_block(load(data(modes=["self_drive"], days=[d1, day2(), day3()])), 1)
    assert "PARKING\nValet available" in block
    assert "PARKING" not in day_block(load(), 1)


def test_drive_day_stops_and_total():
    t = trip()
    t.days[1].drive_stops = [Stop(name="Sile", minutes_from_prev=75, suggested_stop_min=60, km_from_prev=70.0),
                             Stop(name="Agva", minutes_from_prev=40, suggested_stop_min=60, km_from_prev=32.5)]
    block = day_block(load(data(trip=t)), 2)
    assert "Sile — 70 km · 1 h 15 min" in block
    assert "Agva — 32.5 km · 40 min" in block
    assert "Total today: 102.5 km · 1 h 55 min" in block


def test_cruise_port_day():
    t = trip(hotels=[HotelStay(city="Mykonos", hotel="MS Aurora", check_in=D1, check_out=D3, source="voucher",
                               kind="cruise")])
    dock = TransportLeg(kind="boat", from_place="Sea", to_place="Mykonos Port", arrive=datetime(2026, 10, 30, 8, 0),
                        source="cruise")
    sail = TransportLeg(kind="boat", from_place="Mykonos Port", to_place="Sea", depart=datetime(2026, 10, 30, 17, 0),
                        source="cruise")
    d = RenderDay(index=2, date=D2, city="Mykonos", text=DayText(title="Mykonos", intro="Port day."),
                  result=DayResult(), meals=[], legs={}, places={}, booking_by_item={}, price_by_restaurant={},
                  arrival=dock, departure=sail, prev_city="Mykonos")
    block = day_block(load(data(trip=t, days=[day1(), d])), 2)
    assert "Docks at Mykonos Port at 08:00" in block
    assert "Back on board by 16:00" in block
    assert "Overnight on board MS Aurora" in block
    assert "→" not in block


def test_end_sections_verify_and_per_country():
    text = all_text(load())
    money = text[text.rindex("Money & Payments\n"):text.rindex("Mobile Connectivity & SIM Guide\n")]
    assert "About Cards.\n⚠ VERIFY: Cards" in money
    safety = text[text.rindex("Safety & Emergency Contacts"):]
    assert safety.index("Turkey") < safety.index("Call 112.") < safety.index("Greece") < safety.index("Call 112 in Greece.")
    doc = load()
    assert [p.text for p in doc.paragraphs if p.style.name == "Heading 2" and p.text in ("Turkey", "Greece")] == [
        "Turkey", "Greece"]


def test_important_places_and_markers():
    text = all_text(load())
    seg = text[text.rindex("Important Places Around Your Stays\n"):text.rindex("Getting Around\n")]
    for s in ("Pera Palace, Istanbul", "Eczane Pera", "Walk · 4 min", "9–6 Mon–Sun", "+90 212 000 0000",
              "⚠ VERIFY: opening hours unknown", "Amerikan Hastanesi", "By car · 12 min",
              "⚠ VERIFY: fewer than 2 groceries"):
        assert s in seg, s


def test_omitted_sections_dropped_from_toc_and_body():
    secs = sections()
    for k in ("money", "connectivity", "health"):
        secs[k] = None
    doc = load(data(sections=secs, domestic=True, trip_scope="domestic"))
    heads = h1s(doc)
    assert "Money & Payments" not in heads and "Health & Vaccination" not in heads
    assert "Mobile Connectivity" not in all_text(doc)
    assert "Tailored Packing List" in heads


def test_thank_you_page():
    doc = load()
    text = all_text(doc)
    seg = text[text.rindex("Thank You"):]
    assert "journey to Istanbul" in seg and "{destination}" not in seg
    assert "+91 86000 15316" in seg
    assert any("bonvoyagebymarina.com" in r.target_ref for r in doc.part.rels.values() if r.is_external)


def test_sunset_time_paris_midsummer():
    t = sunset_time(48.8566, 2.3522, date(2026, 6, 21), 2.0)
    assert abs((t.hour * 60 + t.minute) - (21 * 60 + 58)) <= 10
    assert sunset_time(78.2, 15.6, date(2026, 6, 21), 2.0) is None


def test_attraction_ai_duration_best_time_tips_marked_unless_library_duration():
    d2 = day2()
    d2.text.attractions[1] = AttractionText(name="Hagia Sophia", description="A great mosque.", duration="1 hour",
                                            best_time="Early morning", tips="Remove shoes")
    d2.text.attractions[0].duration = "3 hours"
    d2.library_attractions = {"Topkapi Palace": {"recommended_duration": "2-3 hours"}}
    block = day_block(load(data(days=[day1(), d2, day3()])), 2)
    hagia = block[block.index("Hagia Sophia\n"):block.index("Bosphorus Sunset Cruise\n")]
    assert "DURATION\n1 hour · ⚠ VERIFY: duration" in hagia
    assert "BEST TIME\nEarly morning · ⚠ VERIFY: best time" in hagia
    assert "TIP\nRemove shoes · ⚠ VERIFY: tip" in hagia
    topkapi = block[block.index("Topkapi Palace\n"):block.index("Hagia Sophia\n")]
    assert "DURATION\n2-3 hours\n" in topkapi
    assert "VERIFY: duration" not in topkapi and "\n3 hours" not in topkapi


def test_intercity_day_without_leg_shows_hotel_link_and_verify_note():
    t = trip(hotels=[HotelStay(city="Paris", hotel="Hotel P", check_in=D1, check_out=D2, source="voucher"),
                     HotelStay(city="Lyon", hotel="Hotel L", check_in=D2, check_out=D3, source="voucher")])
    d = RenderDay(index=2, date=D2, city="Lyon", text=DayText(title="On to Lyon", intro="Driving today."),
                  result=DayResult(), places={"Hotel L": place("Hotel L", note="address differs from voucher")},
                  prev_city="Paris")
    doc = load(data(trip=t, days=[day1(), d]))
    block = day_block(doc, 2)
    assert "Paris → Hotel L" in block
    assert "⚠ VERIFY: transfer time" in block and "⚠ VERIFY: address differs from voucher" in block
    card = next(t for t in doc.tables if t.cell(0, 0).paragraphs[0].text == "Paris → Hotel L")
    rid = card.cell(0, 0).paragraphs[0]._p.find(qn("w:hyperlink")).get(qn("r:id"))
    assert "query=Hotel+L" in doc.part.rels[rid].target_ref


def test_new_stay_without_arrival_leg_gets_check_in_card():
    t = trip(hotels=[HotelStay(city="Istanbul", hotel="Hotel A", check_in=D1, check_out=D2, source="voucher"),
                     HotelStay(city="Istanbul", hotel="Hotel B", check_in=D2, check_out=D3, source="voucher",
                               check_in_time=time(14))])
    d = RenderDay(index=2, date=D2, city="Istanbul", text=DayText(title="Move", intro="New hotel."),
                  result=DayResult(), places={"Hotel B": None}, prev_city="Istanbul")
    doc = load(data(trip=t, days=[day1(), d]))
    block = day_block(doc, 2)
    assert "Hotel B\nCHECK-IN\nfrom 14:00" in block
    assert "⚠ VERIFY: Hotel B not found on Google" in block
    card = next(t for t in doc.tables if t.cell(0, 0).paragraphs[0].text == "Hotel B")
    assert card.cell(0, 0).paragraphs[0]._p.find(qn("w:hyperlink")) is not None


def test_anchored_slot_without_pick_marks_each_anchor():
    d1 = day1()
    d1.meals = []
    d1.result.meals[0].anchors = [("Galata Tower", (41.0, 29.0), 1), (HOTEL, (41.0, 28.9), 2)]
    block = day_block(load(data(days=[d1, day2(), day3()])), 1)
    assert block.count("⚠ ADD RESTAURANT – Dinner, Day 1, near Galata Tower") == 1
    assert block.count(f"⚠ ADD RESTAURANT – Dinner, Day 1, near {HOTEL}") == 2
    assert "near Istanbul" not in block


def test_pick_not_reused_for_other_slot_of_same_meal():
    d1 = day1()
    d1.meals[0].slot = MealSlot(meal="dinner", start=time(21), anchors=d1.result.meals[0].anchors)
    block = day_block(load(data(days=[d1, day2(), day3()])), 1)
    assert "Mikla" not in block
    assert block.count(f"⚠ ADD RESTAURANT – Dinner, Day 1, near {HOTEL}") == 3


def test_dear_guest_when_no_client_name_and_transport_card_without_place_has_no_link():
    t = trip(client_names=[])
    boat = TransportLeg(kind="boat", from_place="", to_place="", arrive=datetime(2026, 10, 29, 10, 0), source="x")
    doc = load(data(trip=t, days=[day1().model_copy(update={"arrival": boat}), day2(), day3()]))
    text = all_text(doc)
    assert "Dear Guest," in text and "⚠ WRITE PERSONAL LETTER" in text
    card = next(t for t in doc.tables if t.cell(0, 0).paragraphs[0].text == "Boat")
    assert card.cell(0, 0).paragraphs[0]._p.find(qn("w:hyperlink")) is None
    assert "query=&" not in doc.element.xml
    assert all(not r.target_ref.endswith("query=") for r in doc.part.rels.values() if r.is_external)


def test_sunset_skips_unmatched_hotel_place():
    d2 = day2()
    d2.places["Bosphorus Sunset Cruise"] = None
    d2.places[HOTEL] = place(HOTEL, match_ok=False)
    block = day_block(load(data(days=[day1(), d2, day3()])), 2)
    assert "SUNSET" not in block
    d2.places[HOTEL] = place(HOTEL)
    assert "SUNSET" in day_block(load(data(days=[day1(), d2, day3()])), 2)


def test_slot_heading_repeats_after_meal():
    d1 = day1()
    noon, later = item("Spice Bazaar", slot="afternoon"), item("Grand Bazaar", slot="afternoon")
    lunch = MealSlot(meal="lunch", start=time(12, 30), included=True)
    d1.result = DayResult(scheduled=[sched(noon, "afternoon", time(12)), sched(later, "afternoon", time(14))],
                          meals=[lunch])
    d1.meals = []
    block = day_block(load(data(days=[d1, day2(), day3()])), 1)
    after = block[block.index("Lunch — Included in your package"):]
    assert "Afternoon\nGrand Bazaar" in after or after.index("Afternoon") < after.index("Grand Bazaar")
