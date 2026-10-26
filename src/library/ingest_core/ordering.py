"""Free pre-pass: scan past AIGs and fix the processing order before extraction.

Order (spec section 6): format rank (old before new), then a trusted date
(file mtime unless it is a bulk-copy date, else the latest travel day found
in weekday-plus-date day headers), then file name. Duplicates and files that
do not look like AIGs (or cannot be read) are kept in the list but marked
skipped and sorted last.
"""

from __future__ import annotations

import hashlib
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Literal

from .text import UnreadableFile, read_text, text_hash

YEARS = range(2023, 2028)
MATCH_RATIO = 0.6
MIN_HEADERS = 2
MIN_OLD_ROWS = 3
MIN_DAY_LINES = 2
MIN_DATED_LINES = 3
HEAD_LINES = 3

# Table first cells (old format): "Day 1", "Day 2", ...
_DAY_RE = re.compile(r"^\s*day\s*\d+", re.IGNORECASE)
# Text lines: same, but new-format headings often start with an emoji
# ("📅 Day 1: May 23 — Arrival"), so leading non-word characters are allowed.
_DAY_LINE_RE = re.compile(r"^[\W_]*day\s*\d+", re.IGNORECASE)

_WEEKDAYS = {
    "mon": 0, "monday": 0,
    "tue": 1, "tues": 1, "tuesday": 1,
    "wed": 2, "weds": 2, "wednesday": 2,
    "thu": 3, "thur": 3, "thurs": 3, "thursday": 3,
    "fri": 4, "friday": 4,
    "sat": 5, "saturday": 5,
    "sun": 6, "sunday": 6,
}
_MONTHS = {
    "jan": 1, "january": 1, "feb": 2, "february": 2, "mar": 3, "march": 3,
    "apr": 4, "april": 4, "may": 5, "jun": 6, "june": 6, "jul": 7, "july": 7,
    "aug": 8, "august": 8, "sep": 9, "sept": 9, "september": 9,
    "oct": 10, "october": 10, "nov": 11, "november": 11, "dec": 12, "december": 12,
}


def _alt(words) -> str:
    return "|".join(sorted(words, key=len, reverse=True))


_WD = rf"(?P<wd>{_alt(_WEEKDAYS)})\.?"
_MO = rf"(?P<mo>{_alt(_MONTHS)})\.?"
_WD_ANY = rf"(?:{_alt(_WEEKDAYS)})\.?"
_MO_ANY = rf"(?:{_alt(_MONTHS)})\.?"
_DD = r"(?P<dd>\d{1,2})(?:st|nd|rd|th)?"
# "Mon, 12 Jan" / "Monday 12th of January"
_DAY_FIRST = re.compile(rf"\b{_WD},?\s+{_DD}\s+(?:of\s+)?{_MO}\b", re.IGNORECASE)
# "Monday, Jan 12" / "Mon Jan 12th"
_MONTH_FIRST = re.compile(rf"\b{_WD},?\s+{_MO}\s+{_DD}\b", re.IGNORECASE)

# Day lines in guides that never write "Day N" (after an optional emoji/bullet):
# a weekday followed by "|", ",", ":" or a date ("Saturday | Arrival",
# "Sat, Nov 30 | ...", "Monday 12 Jan ..."), or
# a month + day number followed by "|", ":" or a dash ("Mar 21 | ...",
# "📅 May 29: ...", "Dec 16, Monday | ..."), or day + month ("21 May – ...").
_WEEKDAY_LINE_RE = re.compile(
    rf"^[\W_]*{_WD_ANY}\b\s*(?:[|,:]|\d{{1,2}}\b|{_MO_ANY}\b)", re.IGNORECASE
)
_MONTH_LINE_RE = re.compile(
    rf"^[\W_]*{_MO_ANY}\s+\d{{1,2}}(?:st|nd|rd|th)?\b(?:,?\s*{_WD_ANY})?\s*[|:\u2013\u2014-]",
    re.IGNORECASE,
)
# Same, date-first: "21 May – Arrival in Cairns".
_DAY_MONTH_LINE_RE = re.compile(
    rf"^[\W_]*\d{{1,2}}(?:st|nd|rd|th)?\s+{_MO_ANY}\b(?:,?\s*{_WD_ANY})?\s*[|:\u2013\u2014-]",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class FileScan:
    path: Path
    mtime: date
    text_hash: str
    fmt: Literal["new", "old"]
    last_day: date | None
    looks_like_aig: bool
    head_lines: tuple[str, ...] = ()
    unreadable: bool = False


@dataclass
class OrderedFile:
    scan: FileScan
    order_date: date | None
    basis: Literal["mtime", "travel", "format", "name"]
    duplicate_of: Path | None
    skipped_reason: str | None


def _headers(text: str) -> list[tuple[int, int, int]]:
    """Unique (weekday, day, month) triples found in the text."""
    found: set[tuple[int, int, int]] = set()
    for rx in (_DAY_FIRST, _MONTH_FIRST):
        for m in rx.finditer(text):
            wd = _WEEKDAYS[m.group("wd").lower()]
            mo = _MONTHS[m.group("mo").lower()]
            dd = int(m.group("dd"))
            if 1 <= dd <= 31:
                found.add((wd, dd, mo))
    return sorted(found, key=lambda h: (h[2], h[1], h[0]))


def _infer_last_day(text: str) -> date | None:
    headers = _headers(text)
    if len(headers) < MIN_HEADERS:
        return None
    best: list[tuple[int, int]] = []  # (matches, year)
    for year in YEARS:
        matches = 0
        for wd, dd, mo in headers:
            try:
                d = date(year, mo, dd)
            except ValueError:
                continue  # e.g. 29 Feb in a non-leap year
            if d.weekday() == wd:
                matches += 1
        if matches >= MIN_HEADERS and matches >= MATCH_RATIO * len(headers):
            best.append((matches, year))
    if not best:
        return None
    best.sort(reverse=True)
    if len(best) > 1 and best[0][0] == best[1][0]:
        return None  # ambiguous between years
    year = best[0][1]
    days = []
    for wd, dd, mo in headers:
        try:
            d = date(year, mo, dd)
        except ValueError:
            continue
        if d.weekday() == wd:
            days.append(d)
    return max(days) if days else None


def _iter_tables(tables):
    for table in tables:
        yield table
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_tables(cell.tables)


def _has_day_table(path: Path) -> bool:
    from docx import Document

    doc = Document(str(path))
    for table in _iter_tables(doc.tables):
        day_rows = 0
        for row in table.rows:
            cells = row.cells
            if cells and _DAY_RE.match(cells[0].text):
                day_rows += 1
        if day_rows >= MIN_OLD_ROWS:
            return True
    return False


def _file_mtime(path: Path) -> date:
    try:
        return date.fromtimestamp(path.stat().st_mtime)
    except (OSError, ValueError, OverflowError):
        return date.min


def _looks_like_aig(lines: list[str]) -> bool:
    if sum(1 for line in lines if _DAY_LINE_RE.match(line)) >= MIN_DAY_LINES:
        return True
    dated = sum(
        1
        for line in lines
        if _WEEKDAY_LINE_RE.match(line)
        or _MONTH_LINE_RE.match(line)
        or _DAY_MONTH_LINE_RE.match(line)
    )
    return dated >= MIN_DATED_LINES


def scan_file(path: Path) -> FileScan:
    path = Path(path)
    mtime = _file_mtime(path)
    try:
        text = read_text(path)
    except UnreadableFile:
        try:
            raw = path.read_bytes()
        except OSError:
            raw = b""
        return FileScan(
            path=path,
            mtime=mtime,
            text_hash=hashlib.sha1(raw).hexdigest(),
            fmt="new",
            last_day=None,
            looks_like_aig=False,
            unreadable=True,
        )

    fmt: Literal["new", "old"] = "new"
    if path.suffix.lower() == ".docx":
        try:
            if _has_day_table(path):
                fmt = "old"
        except Exception:
            fmt = "new"

    lines = text.splitlines()
    head = tuple(line.strip() for line in lines if line.strip())[:HEAD_LINES]
    return FileScan(
        path=path,
        mtime=mtime,
        text_hash=text_hash(text),
        fmt=fmt,
        last_day=_infer_last_day(text),
        looks_like_aig=_looks_like_aig(lines),
        head_lines=head,
    )


def _fmt_rank(scan: FileScan) -> int:
    return 0 if scan.fmt == "old" else 1


def _name_key(scan: FileScan) -> tuple[str, str]:
    return (scan.path.name, str(scan.path))


def order_files(scans: list[FileScan], bulk_threshold: int = 5) -> list[OrderedFile]:
    """Fix the processing order (oldest first; later files overwrite earlier).

    Sort key: processable before skipped, format rank (old=0, new=1),
    ``order_date or date.min``, file name. Undated files therefore sort before
    dated files of the same format, so dated data overwrites undated data: a
    conservative choice, since undated files have no trusted date.

    Skip precedence: ``unreadable``, then ``not_an_aig``, then ``duplicate``.
    Duplicates are detected only among processable files; the first by file
    name is kept and the rest get ``duplicate_of``.
    """
    mtime_counts = Counter(s.mtime for s in scans)

    def _base_reason(s: FileScan) -> str | None:
        if s.unreadable:
            return "unreadable"
        if not s.looks_like_aig:
            return "not_an_aig"
        return None

    # Duplicates: first by path name keeps the text, the rest are skipped.
    by_hash: dict[str, list[FileScan]] = defaultdict(list)
    for s in scans:
        if _base_reason(s) is None:
            by_hash[s.text_hash].append(s)
    dup_of: dict[int, Path] = {}
    for group in by_hash.values():
        group = sorted(group, key=_name_key)
        for s in group[1:]:
            dup_of[id(s)] = group[0].path

    entries: list[OrderedFile] = []
    for s in scans:
        if mtime_counts[s.mtime] <= bulk_threshold:
            order_date, basis = s.mtime, "mtime"
        elif s.last_day is not None:
            order_date, basis = s.last_day, "travel"
        else:
            order_date, basis = None, "format"  # refined below
        reason = _base_reason(s)
        if reason is None and id(s) in dup_of:
            reason = "duplicate"
        entries.append(
            OrderedFile(
                scan=s,
                order_date=order_date,
                basis=basis,
                duplicate_of=dup_of.get(id(s)),
                skipped_reason=reason,
            )
        )

    # Undated files that tie with another undated processable file of the
    # same format are ordered by name, not by format.
    undated = Counter(
        _fmt_rank(e.scan)
        for e in entries
        if e.order_date is None and e.skipped_reason is None
    )
    for e in entries:
        if e.order_date is None and e.skipped_reason is None and undated[_fmt_rank(e.scan)] > 1:
            e.basis = "name"

    entries.sort(
        key=lambda e: (
            0 if e.skipped_reason is None else 1,
            _fmt_rank(e.scan),
            e.order_date or date.min,
            *_name_key(e.scan),
        )
    )
    return entries


def _md(value: str) -> str:
    return value.replace("|", "\\|")


def _reason_text(e: OrderedFile) -> str:
    if e.skipped_reason == "duplicate" and e.duplicate_of is not None:
        return f"duplicate of {e.duplicate_of.name}"
    return e.skipped_reason or ""


def write_order_report(ordered: list[OrderedFile], out: Path) -> None:
    out = Path(out)
    processed = [e for e in ordered if e.skipped_reason is None]
    lines = [
        "# Ingest processing order",
        "",
        f"{len(ordered)} files scanned; {len(processed)} to process, "
        f"{len(ordered) - len(processed)} skipped.",
        "",
        "| position | file | order_date | basis | format | skipped reason |",
        "|---|---|---|---|---|---|",
    ]
    pos = 0
    for e in ordered:
        if e.skipped_reason is None:
            pos += 1
            position = str(pos)
        else:
            position = "-"
        reason = _md(_reason_text(e))
        name = _md(e.scan.path.name)
        order_date = e.order_date.isoformat() if e.order_date else ""
        basis = e.basis if e.skipped_reason is None else ""
        lines.append(
            f"| {position} | {name} | {order_date} | {basis} | {e.scan.fmt} | {reason} |"
        )

    fallback = [e for e in processed if e.basis in ("format", "name")]
    lines += [
        "",
        "## Check before the paid run: ordered by format or file name",
        "",
        "These files have no trusted date; their position comes from format "
        "(old before new) and then file name. Correct the order if needed.",
        "",
    ]
    if fallback:
        lines += [f"- {e.scan.path.name} ({e.basis}, {e.scan.fmt})" for e in fallback]
    else:
        lines.append("- none")

    skipped = [e for e in ordered if e.skipped_reason is not None]
    lines += [
        "",
        "## Skipped files (check these)",
        "",
        "Rescue any real guide skipped by mistake. For non-guides and unreadable "
        "files the first lines of text are shown.",
        "",
    ]
    if not skipped:
        lines.append("- none")
    for e in skipped:
        lines.append(f"- {e.scan.path.name}: {_reason_text(e)}")
        if e.skipped_reason in ("not_an_aig", "unreadable"):
            if e.scan.unreadable:
                lines.append("    - (unreadable)")
            elif e.scan.head_lines:
                lines += [f"    - {line}" for line in e.scan.head_lines]
            else:
                lines.append("    - (no text)")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
