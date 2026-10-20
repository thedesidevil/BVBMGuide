import io
import zipfile
from pathlib import Path

import fitz
import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image

from src.aigv2.intake.attachments import read_attachment, read_attachments


def _png_bytes() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (40, 40), (200, 30, 30)).save(buf, "PNG")
    return buf.getvalue()


def _text_pdf(path: Path, text: str) -> Path:
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), text)
    doc.save(path)
    return path


def _image_pdf(path: Path, text: str = "") -> Path:
    doc = fitz.open()
    page = doc.new_page()
    if text:
        page.insert_text((72, 72), text)
    page.insert_image(fitz.Rect(72, 150, 272, 350), stream=_png_bytes())
    doc.save(path)
    return path


def test_docx_text_and_tables(tmp_path):
    d = Document()
    d.add_paragraph("Welcome paragraph one")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Cell Alpha"
    t.cell(0, 1).text = "Cell Beta"
    d.add_paragraph("Closing paragraph")
    p = tmp_path / "a.docx"
    d.save(p)
    a = read_attachment(p)
    assert a.kind == "docx"
    assert a.name == "a.docx"
    for s in ("Welcome paragraph one", "Cell Alpha", "Cell Beta", "Closing paragraph"):
        assert s in a.text
    assert a.text.index("Welcome") < a.text.index("Cell Alpha") < a.text.index("Closing")
    assert a.has_images is False and a.image_only is False


def test_docx_with_inline_picture(tmp_path):
    png = tmp_path / "pic.png"
    png.write_bytes(_png_bytes())
    d = Document()
    d.add_paragraph("3 senior citizen women")
    d.add_paragraph().add_run().add_picture(str(png), width=Inches(1))
    p = tmp_path / "notes.docx"
    d.save(p)
    a = read_attachment(p)
    assert a.has_images is True
    assert "3 senior citizen women" in a.text
    assert "Hagia Sophia" not in a.text
    assert a.image_only is True or len(a.text.strip()) >= 40


def test_docx_image_only(tmp_path):
    png = tmp_path / "pic.png"
    png.write_bytes(_png_bytes())
    d = Document()
    d.add_paragraph().add_run().add_picture(str(png), width=Inches(1))
    p = tmp_path / "scan.docx"
    d.save(p)
    a = read_attachment(p)
    assert a.has_images and a.image_only


def test_pdf_with_text(tmp_path):
    p = _text_pdf(tmp_path / "t.pdf", "Day 1 arrival in Istanbul, hotel check in 14 Oct 2026 at the Grand Hotel")
    a = read_attachment(p)
    assert a.kind == "pdf"
    assert "Istanbul" in a.text
    assert a.image_only is False and a.has_images is False
    assert "14 Oct 2026" in a.text_dates


def test_pdf_image_only(tmp_path):
    a = read_attachment(_image_pdf(tmp_path / "scan.pdf"))
    assert a.kind == "pdf"
    assert a.image_only is True and a.has_images is True
    assert a.text.strip() == ""


def test_pdf_boilerplate_plus_image_no_date(tmp_path):
    p = _image_pdf(
        tmp_path / "voucher.pdf",
        "Booking confirmation. Please present this voucher at check in. Terms and conditions apply.",
    )
    a = read_attachment(p)
    assert a.has_images is True
    assert a.image_only is False
    assert a.text_dates == []


@pytest.mark.parametrize(
    "text,expected",
    [
        ("Arrive 14 Oct 2026 late", "14 Oct 2026"),
        ("Arrive 14/10/2026 late", "14/10/2026"),
        ("Arrive 2026-10-14 late", "2026-10-14"),
        ("Arrive Oct 14 late", "Oct 14"),
        ("Arrive 14 October late", "14 October"),
    ],
)
def test_text_dates_formats(tmp_path, text, expected):
    d = Document()
    d.add_paragraph(text + " and enough filler text to pass the threshold")
    p = tmp_path / "d.docx"
    d.save(p)
    assert expected in read_attachment(p).text_dates


@pytest.mark.parametrize("ext", [".png", ".jpg", ".jpeg", ".heic", ".webp"])
def test_image_files(tmp_path, ext):
    p = tmp_path / f"x{ext}"
    p.write_bytes(_png_bytes())
    a = read_attachment(p)
    assert a.kind == "image" and a.image_only is True and a.text == ""


def test_other_kind(tmp_path):
    p = tmp_path / "x.txt"
    p.write_text("hello")
    a = read_attachment(p)
    assert a.kind == "other" and a.text == ""


def test_zip_expansion_and_malicious_entry(tmp_path):
    inner = _text_pdf(tmp_path / "i.pdf", "Voucher for Hotel Sultan, 14 Oct 2026 check in, two nights")
    z = tmp_path / "bundle.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.write(inner, "docs/voucher.pdf")
        zf.writestr("../evil.pdf", inner.read_bytes())
        zf.writestr("/abs.pdf", inner.read_bytes())
    atts = read_attachments([z])
    assert [a.name for a in atts] == ["bundle.zip/docs/voucher.pdf"]
    assert "Hotel Sultan" in atts[0].text


def test_zip_caps(tmp_path):
    z = tmp_path / "many.zip"
    with zipfile.ZipFile(z, "w") as zf:
        for i in range(60):
            zf.writestr(f"f{i}.txt", "x")
    assert len(read_attachments([z])) == 50

    big = tmp_path / "big.zip"
    with zipfile.ZipFile(big, "w") as zf:
        zf.writestr("a.txt", b"0" * (30 * 1024 * 1024))
        zf.writestr("b.txt", b"0" * (30 * 1024 * 1024))
    assert [a.name for a in read_attachments([big])] == ["big.zip/a.txt"]


def test_dedup_by_hash(tmp_path):
    a = _text_pdf(tmp_path / "one.pdf", "Same content voucher number 12345 for the hotel stay")
    b = tmp_path / "two.pdf"
    b.write_bytes(a.read_bytes())
    z = tmp_path / "z.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.write(a, "copy.pdf")
    atts = read_attachments([a, b, z])
    assert len(atts) == 3
    assert atts[0].duplicate_of is None
    assert atts[1].duplicate_of == "one.pdf"
    assert atts[2].duplicate_of == "one.pdf"
