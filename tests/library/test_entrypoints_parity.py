"""CLI (LibraryBuilder / run_ingest) and UI (IngestService) share the ingest core."""

import os
from datetime import date, datetime
from pathlib import Path

import pytest
from docx import Document

from src.library.ingest_core import PlanEntry, run_ingest
from src.library.ingest_core import extract as extract_mod
from src.library.ingest_core import runner as runner_mod
from src.library.ingest_core.extract import Extraction
from src.library.ingest_core.ordering import order_files, scan_file
from src.library.ui.services.db_service import LibraryDBService
from src.library.ui.services.ingest_service import IngestService
from src.library.ui.storage import LocalStorageBackend

from .conftest import MemoryBackend

LAST_DAY = date(2026, 1, 14)
DAY_LINES = [
    "Sakura Adventures: Italy",
    "Day 1: Monday, 12 Jan – Arrival in Rome",
    "Day 2: Tuesday, 13 Jan – Vatican",
    "Day 3: Wednesday, 14 Jan – Departure",
]


class StubClient:
    def __init__(self):
        self.cost_usd = 0.0
        self.usage = {"prompt_tokens": 0, "completion_tokens": 0}


def _set_mtime(path: Path, d: date) -> None:
    ts = datetime(d.year, d.month, d.day, 12, 0, 0).timestamp()
    os.utime(path, (ts, ts))


def _doc(path: Path, extra: str = "", mtime: date | None = LAST_DAY) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for line in DAY_LINES:
        doc.add_paragraph(line)
    if extra:
        doc.add_paragraph(extra)
    doc.save(path)
    if mtime:
        _set_mtime(path, mtime)
    return path


def _extraction_data(hours="12-15"):
    return {
        "covered_cities": ["Rome"],
        "city_countries": {"Rome": "Italy"},
        "restaurants": [{"name": "Roscioli", "city": "Rome", "hours": hours}],
        "attractions": [{"name": "Colosseum", "city": "Rome", "entry_fee": "18 EUR"}],
        "safety_tips": [{"cities": ["Rome", "Florence"], "tip": "Watch for pickpockets on bus 64"}],
        "country_content": [{"country": "Italy", "currency": "EUR",
                             "tipping": [{"tip": "Rounding up is enough"}]}],
    }


def _shards(backend) -> dict:
    out = {}
    for key in backend.list_keys(""):
        if key == "_index.json" or key.startswith("_staging/") or key.startswith("_logs/"):
            continue
        out[key] = backend.read_json(key)
    return out


def _ui_persist(backend, file_path: Path, folder: str, data: dict, library_path=None) -> dict:
    service = IngestService(backend, library_path)
    session = service.create_session()
    [entry] = service.save_upload(session, file_path.name, file_path.read_bytes())
    session.update_file(entry["file_id"], {"state": "extracted", "assigned_folder": folder})
    service.save_extracted_edits(session, entry["file_id"], data)
    return service.persist(session)


# --- parity -------------------------------------------------------------

def test_ui_persist_and_cli_run_ingest_produce_identical_shards(tmp_path, monkeypatch):
    path = _doc(tmp_path / "lib" / "Italy" / "trip.docx")
    data = _extraction_data()

    # CLI path
    cli_backend = MemoryBackend()
    cli_db = LibraryDBService(cli_backend)
    [ordered] = order_files([scan_file(path)])
    assert ordered.skipped_reason is None
    plan = [PlanEntry(path=path, folder="Italy", order_date=ordered.order_date,
                      fmt=ordered.scan.fmt, text_hash=ordered.scan.text_hash)]
    monkeypatch.setattr(runner_mod, "extract_file",
                        lambda client, p: Extraction(data=data, chunks=1, prompt_tokens=0,
                                                     completion_tokens=0, usd=0.0))
    report = run_ingest(plan, cli_db, StubClient(), log_dir=tmp_path / "logs", workers=1)
    assert report.processed == ["Italy/trip.docx"]

    # UI path
    ui_backend = MemoryBackend()
    result = _ui_persist(ui_backend, path, "Italy", data)

    assert _shards(ui_backend) == _shards(cli_backend)
    assert "Rome.json" in _shards(ui_backend)
    assert "_country/Italy.json" in _shards(ui_backend)
    ui_index = ui_backend.read_json("_index.json")
    cli_index = cli_backend.read_json("_index.json")
    assert ui_index["_folder_coverage"] == cli_index["_folder_coverage"] == {"Italy": ["Rome"]}
    assert set(ui_index["_processed_files"]) == set(cli_index["_processed_files"]) == {"Italy/trip.docx"}
    ui_rec = ui_index["_processed_files"]["Italy/trip.docx"]
    cli_rec = cli_index["_processed_files"]["Italy/trip.docx"]
    for k in ("text_hash", "order_date", "fmt", "folder", "covered_cities"):
        assert ui_rec[k] == cli_rec[k]

    assert result["persisted_files"] == 1
    assert result["affected_cities"] == ["Florence", "Rome"]
    assert result["stats"]["added"] == report.stats.added
    assert "Rome" in result["stats"]["by_city"]
    assert result["dropped"] == 0
    # a re-run of the CLI over the file the UI already persisted skips it (same hash)
    ui_db = LibraryDBService(ui_backend)
    again = run_ingest(plan, ui_db, StubClient(), log_dir=tmp_path / "logs2", workers=1)
    assert again.skipped == [("Italy/trip.docx", "already_processed")]


def test_ui_persist_replaces_older_same_name_item(tmp_path, memory_backend):
    db = LibraryDBService(memory_backend)
    db.save_city_data("Rome", {
        "restaurants": [{"name": "Roscioli", "city": "Rome", "hours": "old hours",
                         "_source_file": "Italy/old.docx", "_source_date": "2024-01-01",
                         "_source_format": "new"}],
        "source_files": ["Italy/old.docx"],
    })
    path = _doc(tmp_path / "trip.docx")
    result = _ui_persist(memory_backend, path, "Italy", _extraction_data(hours="new hours"))

    [item] = memory_backend.read_json("Rome.json")["restaurants"]
    assert item["hours"] == "new hours"  # old UI code skipped it as a duplicate
    assert item["_replaced_from"] == "Italy/old.docx"
    assert item["_source_file"] == "Italy/trip.docx"
    assert item["_source_date"] == LAST_DAY.isoformat()
    assert result["stats"]["replaced"] == 1


def test_ui_persist_order_date_falls_back_to_today(tmp_path, memory_backend):
    path = tmp_path / "undated.docx"
    doc = Document()
    doc.add_paragraph("A guide without dated day headers")
    doc.save(path)
    _ui_persist(memory_backend, path, "Italy", _extraction_data())
    [item] = memory_backend.read_json("Rome.json")["restaurants"]
    assert item["_source_date"] == date.today().isoformat()


def test_ui_persist_cleans_and_counts_dropped(tmp_path, memory_backend):
    data = _extraction_data()
    data["restaurants"].append({"name": "ADD RESTAURANT", "city": "Rome"})
    path = _doc(tmp_path / "trip.docx")
    result = _ui_persist(memory_backend, path, "Italy", data)
    names = [r["name"] for r in memory_backend.read_json("Rome.json")["restaurants"]]
    assert names == ["Roscioli"]
    assert result["dropped"] == 1


def test_ui_persist_skips_excluded_and_unassigned_files(tmp_path, memory_backend):
    lib = tmp_path / "lib"
    service = IngestService(memory_backend, lib)
    session = service.create_session()
    ids = []
    for name in ("a.docx", "b.docx", "c.docx"):
        path = _doc(tmp_path / name, extra=name)
        [entry] = service.save_upload(session, name, path.read_bytes())
        ids.append(entry["file_id"])
        service.save_extracted_edits(session, entry["file_id"], _extraction_data())
    session.update_file(ids[0], {"assigned_folder": "Italy", "excluded": True})
    session.update_file(ids[1], {"assigned_folder": None})
    session.update_file(ids[2], {"assigned_folder": "Italy"})

    result = service.persist(session)

    assert result["persisted_files"] == 1
    index = memory_backend.read_json("_index.json")
    assert set(index["_processed_files"]) == {"Italy/c.docx"}
    assert sorted(p.name for p in (lib / "Italy").iterdir()) == ["c.docx"]
    assert not memory_backend.list_keys("_staging/")


def test_ui_extract_all_stores_extraction_data(tmp_path, memory_backend, monkeypatch):
    service = IngestService(memory_backend, tmp_path / "lib")
    session = service.create_session()
    good = _doc(tmp_path / "good.docx")
    bad = _doc(tmp_path / "bad.docx", extra="bad")
    [g] = service.save_upload(session, "good.docx", good.read_bytes())
    [b] = service.save_upload(session, "bad.docx", bad.read_bytes())
    data = _extraction_data()
    sizes = {len(good.read_bytes()): data, len(bad.read_bytes()): None}

    def fake_extract(client, path):
        d = sizes[len(Path(path).read_bytes())]
        return None if d is None else Extraction(data=d, chunks=1, prompt_tokens=0,
                                                 completion_tokens=0, usd=None)

    monkeypatch.setattr(extract_mod, "extract_file", fake_extract)
    monkeypatch.setattr("src.common.ai_provider.get_ai_client", lambda **kw: StubClient())

    results = service.extract_all(session, workers=2)

    assert results[g["file_id"]] == "extracted"
    assert memory_backend.read_json(session.get_extracted_key(g["file_id"])) == data
    assert session.get_file(g["file_id"])["state"] == "extracted"
    assert session.get_file(b["file_id"])["state"] == "failed"


# --- LibraryBuilder.build_database --------------------------------------

@pytest.fixture
def builder_env(tmp_path, monkeypatch):
    import src.library.builder as builder_mod

    monkeypatch.setattr(builder_mod, "get_ai_client", lambda **kw: StubClient())
    lib = tmp_path / "aig-library"
    _doc(lib / "Italy" / "rome.docx", extra="rome", mtime=date(2025, 3, 1))
    _doc(lib / "France" / "paris.docx", extra="paris", mtime=date(2025, 4, 1))
    results = {
        "rome.docx": {"covered_cities": ["Rome"],
                      "restaurants": [{"name": "Roscioli", "city": "Rome"}]},
        "paris.docx": {"covered_cities": ["Paris"],
                       "attractions": [{"name": "Louvre", "city": "Paris"}]},
    }
    calls: list[str] = []

    def fake_extract(client, path):
        calls.append(Path(path).name)
        return Extraction(data=results[Path(path).name], chunks=1, prompt_tokens=0,
                          completion_tokens=0, usd=0.0)

    monkeypatch.setattr(runner_mod, "extract_file", fake_extract)
    return builder_mod.LibraryBuilder(lib), tmp_path / "db", calls


def test_build_database_over_tiny_library(builder_env):
    builder, out, calls = builder_env
    result = builder.build_database(output_path=out, workers=1)

    assert sorted(calls) == ["paris.docx", "rome.docx"]
    assert (out / "Rome.json").exists() and (out / "Paris.json").exists()
    assert set(result["destinations"]) == {"Rome", "Paris"}
    assert result["destinations"]["Rome"]["restaurants"][0]["name"] == "Roscioli"
    assert set(result["_processed_files"]) == {"Italy/rome.docx", "France/paris.docx"}
    assert result["_folder_coverage"] == {"Italy": ["Rome"], "France": ["Paris"]}
    index = LocalStorageBackend(out).read_json("_index.json")
    assert index["_folder_coverage"] == result["_folder_coverage"]

    # second run: nothing changed, nothing re-extracted
    calls.clear()
    builder.build_database(output_path=out, workers=1)
    assert calls == []


def test_build_database_force_reextracts_without_deleting_shards(builder_env):
    builder, out, calls = builder_env
    builder.build_database(output_path=out, workers=1)
    # a shard not produced by these files must survive --force
    LocalStorageBackend(out).write_json("Tokyo.json", {"restaurants": [{"name": "Sushi"}]})

    calls.clear()
    result = builder.build_database(output_path=out, force=True, workers=1)

    assert sorted(calls) == ["paris.docx", "rome.docx"]
    assert (out / "Tokyo.json").exists()
    assert set(result["destinations"]) == {"Rome", "Paris", "Tokyo"}
    assert len(result["destinations"]["Rome"]["restaurants"]) == 1  # idempotent merge


def test_build_database_rejects_json_output(builder_env, tmp_path):
    builder, _, _ = builder_env
    with pytest.raises(ValueError, match="single-file library output is no longer supported"):
        builder.build_database(output_path=tmp_path / "library_db.json")


def test_ui_persist_skips_failed_extraction(tmp_path, memory_backend):
    path = _doc(tmp_path / "trip.docx")
    result = _ui_persist(memory_backend, path, "Italy", {"error": "AI call failed"})
    assert result["persisted_files"] == 0
    index = memory_backend.read_json("_index.json") or {}
    assert not index.get("_processed_files")


# --- fix round 1: persist isolation, v2 skip, missing bytes ----------------

def _v2_doc(path: Path) -> Path:
    import zipfile
    _doc(path)
    with zipfile.ZipFile(path, "a") as z:
        z.writestr("docProps/custom.xml",
                   '<?xml version="1.0"?><Properties xmlns="http://schemas.openxmlformats.org/'
                   'officeDocument/2006/custom-properties"><property fmtid="{x}" pid="2" '
                   'name="bvbm_run_id"><lpwstr>abc</lpwstr></property></Properties>')
    return path


def test_ui_persist_skips_v2_generated_file(tmp_path, memory_backend):
    path = _v2_doc(tmp_path / "v2.docx")
    result = _ui_persist(memory_backend, path, "Italy", _extraction_data())
    assert result["persisted_files"] == 0
    assert result["skipped"] == [("v2.docx", "v2_generated")]
    assert memory_backend.read_json("Rome.json") is None
    assert not (memory_backend.read_json("_index.json") or {}).get("_processed_files")


def _session_with(service, tmp_path, names):
    session = service.create_session()
    ids = []
    for name in names:
        path = _doc(tmp_path / name, extra=name)
        [entry] = service.save_upload(session, name, path.read_bytes())
        session.update_file(entry["file_id"], {"assigned_folder": "Italy"})
        data = _extraction_data()
        data["restaurants"] = [{"name": f"R-{name}", "city": "Rome"}]
        service.save_extracted_edits(session, entry["file_id"], data)
        ids.append(entry["file_id"])
    return session, ids


def test_ui_persist_isolates_a_failing_file(tmp_path, memory_backend, monkeypatch):
    import src.library.ingest_core.merge as merge_mod

    real = merge_mod.merge_extraction

    def flaky(db, data, meta):
        if meta.source_file == "Italy/b.docx":
            raise RuntimeError("boom")
        return real(db, data, meta)

    monkeypatch.setattr(merge_mod, "merge_extraction", flaky)
    service = IngestService(memory_backend, None)
    session, _ = _session_with(service, tmp_path, ["a.docx", "b.docx", "c.docx"])

    result = service.persist(session)

    assert result["persisted_files"] == 2
    assert result["failed"] == [("b.docx", "RuntimeError: boom")]
    index = memory_backend.read_json("_index.json")
    assert set(index["_processed_files"]) == {"Italy/a.docx", "Italy/c.docx"}
    names = {r["name"] for r in memory_backend.read_json("Rome.json")["restaurants"]}
    assert names == {"R-a.docx", "R-c.docx"}


def test_ui_persist_missing_upload_bytes_is_failed(tmp_path, memory_backend):
    service = IngestService(memory_backend, None)
    session, ids = _session_with(service, tmp_path, ["a.docx", "b.docx"])
    memory_backend.delete(session.get_upload_key(ids[0]))

    result = service.persist(session)

    assert result["persisted_files"] == 1
    assert [f for f, _ in result["failed"]] == ["a.docx"]
    assert set(memory_backend.read_json("_index.json")["_processed_files"]) == {"Italy/b.docx"}


# --- fix round 1: build cost guards ----------------------------------------

def test_build_database_refuses_legacy_index(builder_env):
    builder, out, calls = builder_env
    LocalStorageBackend(out).write_json("_index.json", {"_processed_files": {
        "Italy/rome.docx": {"mtime": 1.0, "destination": "Italy", "covered_cities": ["Rome"]}}})
    for force in (False, True):
        with pytest.raises(ValueError, match="built by the old ingester"):
            builder.build_database(output_path=out, workers=1, force=force)
    assert calls == []


def _six_file_library(lib: Path) -> None:
    for i in range(4):
        _doc(lib / "Spain" / f"s{i}.docx", extra=f"spain {i}", mtime=date(2025, 5, i + 1))


def test_build_database_requires_max_usd_for_more_than_five_files(builder_env, monkeypatch):
    builder, out, calls = builder_env
    _six_file_library(builder.library_path)
    with pytest.raises(ValueError, match="--max-usd is required"):
        builder.build_database(output_path=out, workers=1)
    assert calls == []


def test_build_database_runs_six_files_with_max_usd(builder_env, monkeypatch):
    builder, out, calls = builder_env
    _six_file_library(builder.library_path)
    fake = runner_mod.extract_file

    def fake_any(client, path):
        if Path(path).name.startswith("s"):
            calls.append(Path(path).name)
            return Extraction(data={"covered_cities": ["Madrid"],
                                    "restaurants": [{"name": Path(path).stem, "city": "Madrid"}]},
                              chunks=1, prompt_tokens=0, completion_tokens=0, usd=0.0)
        return fake(client, path)

    monkeypatch.setattr(runner_mod, "extract_file", fake_any)
    result = builder.build_database(output_path=out, workers=1, max_usd=5.0)
    assert len(calls) == 6
    assert len(result["_processed_files"]) == 6
    # --force makes the whole plan count again: refused without --max-usd
    with pytest.raises(ValueError, match="--max-usd is required"):
        builder.build_database(output_path=out, workers=1, force=True)


# --- fix round 1: country API is read-only ---------------------------------

def test_country_api_missing_shard_is_404_and_writes_nothing(tmp_path):
    from fastapi.testclient import TestClient

    from src.library.ui import create_app

    db = tmp_path / "db"
    backend = LocalStorageBackend(db)
    backend.write_json("_index.json", {"_folder_coverage": {"Italy": ["Rome"]}})
    backend.write_json("Rome.json", {"safety_tips": [{"tip": "Watch your bag", "cities": ["Rome"]}]})
    before = sorted(backend.list_keys(""))

    client = TestClient(create_app(db_path=db))
    resp = client.get("/api/country/Italy")

    assert resp.status_code == 404
    assert sorted(backend.list_keys("")) == before


def test_clear_processed_files_keeps_rest_of_index(memory_db):
    memory_db.record_processed_file("Italy/a.docx", {"text_hash": "h"})
    memory_db.add_folder_coverage("Italy", ["Rome"])
    memory_db.clear_processed_files()
    assert memory_db._index["_processed_files"] == {}
    assert memory_db.get_folder_coverage() == {"Italy": ["Rome"]}


# --- ExtractionError reasons reach every caller --------------------------

def test_ui_extract_all_reports_extraction_error_reason(tmp_path, memory_backend, monkeypatch):
    from src.library.ingest_core.extract import ExtractionError

    service = IngestService(memory_backend, tmp_path / "lib")
    session = service.create_session()
    [f] = service.save_upload(session, "a.docx", _doc(tmp_path / "a.docx").read_bytes())

    def boom(client, path):
        raise ExtractionError("truncated: a.docx hit the output limit")

    monkeypatch.setattr(extract_mod, "extract_file", boom)
    seen = {}
    monkeypatch.setattr("src.common.ai_provider.get_ai_client",
                        lambda **kw: seen.update(kw) or StubClient())
    results = service.extract_all(session, workers=1)

    assert results[f["file_id"]] == "failed"
    assert session.get_file(f["file_id"])["state"] == "failed"
    stored = memory_backend.read_json(session.get_extracted_key(f["file_id"]))
    assert "truncated" in stored["error"]
    assert seen["timeout"] == 600.0 and seen["max_retries"] == 1  # ingest client settings


def test_builder_failure_reason_reaches_report_and_log(builder_env, monkeypatch):
    from src.library.ingest_core.extract import ExtractionError

    builder, out, calls = builder_env

    def boom(client, path):
        raise ExtractionError("unparseable")

    monkeypatch.setattr(runner_mod, "extract_file", boom)
    builder.build_database(output_path=out, workers=1)
    log = (out / "_logs" / "ingest_log.jsonl").read_text()
    assert "ExtractionError: unparseable" in log


def test_builder_client_uses_ingest_settings(tmp_path, monkeypatch):
    import src.library.builder as builder_mod

    seen = {}
    monkeypatch.setattr(builder_mod, "get_ai_client", lambda **kw: seen.update(kw) or StubClient())
    builder_mod.LibraryBuilder(tmp_path)
    assert seen["timeout"] == 600.0 and seen["max_retries"] == 1


def test_inspect_cli_reports_extraction_error(tmp_path, monkeypatch):
    from typer.testing import CliRunner

    import src.library.__main__ as cli
    import src.library.builder as builder_mod
    from src.library.ingest_core.extract import ExtractionError

    monkeypatch.setattr(builder_mod, "get_ai_client", lambda **kw: StubClient())

    def boom(client, path):
        raise ExtractionError("truncated: x hit the output limit")

    monkeypatch.setattr(extract_mod, "extract_file", boom)
    path = _doc(tmp_path / "a.docx")
    res = CliRunner().invoke(cli.app, ["inspect", str(path), "--library", str(tmp_path)])
    assert res.exit_code == 1
    assert "truncated" in " ".join(res.output.split())


def test_cli_and_ui_record_identical_country_coverage_for_two_country_guide(tmp_path, monkeypatch):
    path = _doc(tmp_path / "lib" / "Europe" / "trip.docx")
    data = {"covered_cities": ["Zurich", "Rome"],
            "city_countries": {"Zurich": "Switzerland", "Rome": "Italy"},
            "restaurants": [{"name": "Kronenhalle", "city": "Zurich"},
                            {"name": "Roscioli", "city": "Rome"}]}
    cli_backend = MemoryBackend()
    [ordered] = order_files([scan_file(path)])
    plan = [PlanEntry(path=path, folder="Europe", order_date=ordered.order_date,
                      fmt=ordered.scan.fmt, text_hash=ordered.scan.text_hash)]
    monkeypatch.setattr(runner_mod, "extract_file",
                        lambda client, p: Extraction(data=data, chunks=1, prompt_tokens=0,
                                                     completion_tokens=0, usd=0.0))
    run_ingest(plan, LibraryDBService(cli_backend), StubClient(), log_dir=tmp_path / "logs", workers=1)
    ui_backend = MemoryBackend()
    _ui_persist(ui_backend, path, "Europe", data)

    expected = {"Europe": ["Zurich", "Rome"], "Switzerland": ["Zurich"], "Italy": ["Rome"]}
    assert cli_backend.read_json("_index.json")["_folder_coverage"] == expected
    assert ui_backend.read_json("_index.json")["_folder_coverage"] == expected
