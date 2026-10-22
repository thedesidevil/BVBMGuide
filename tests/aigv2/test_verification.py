import io

from openpyxl import load_workbook

from src.aigv2.planning.meals import Exclusion
from src.aigv2.qc.facts import FactRow
from src.aigv2.verification import build_sheet

PLACES_HEADER = ["Day", "Place", "Guide says: hours", "Google says: hours", "Guide says: travel",
                 "Google says: travel", "Maps link", "Match", "Checked by"]


def _rows():
    return [
        FactRow(day=1, venue="Cafe Alpha", guide_hours="09:00-17:00", google_hours="Mon 09:00-17:00",
                guide_travel="Walk · 5 min", google_travel="Walk · 6 min",
                link="https://www.google.com/maps/search/?api=1&query=Cafe+Alpha+Istanbul", match=True),
        FactRow(day=2, venue="Bistro Beta", guide_hours="", google_hours="Not verified",
                guide_travel="", google_travel="Not verified",
                link="https://www.google.com/maps/search/?api=1&query=Bistro+Beta+Istanbul", match=False),
    ]


def _exclusions():
    return [Exclusion(name="Curry House", reason="diet", detail="Not vegetarian friendly")]


def _load(data: bytes):
    return load_workbook(io.BytesIO(data))


def test_places_sheet_has_header_and_one_row_per_fact():
    wb = _load(build_sheet(_rows(), _exclusions()))
    ws = wb["Places"]

    assert [c.value for c in ws[1]] == PLACES_HEADER
    assert ws.max_row == 3
    assert ws.freeze_panes == "A2"


def test_places_sheet_maps_fields_and_match_glyphs():
    ws = _load(build_sheet(_rows(), _exclusions()))["Places"]

    assert [c.value for c in ws[2]][:8] == [
        1, "Cafe Alpha", "09:00-17:00", "Mon 09:00-17:00", "Walk · 5 min", "Walk · 6 min",
        "https://www.google.com/maps/search/?api=1&query=Cafe+Alpha+Istanbul", "✓"]
    assert ws[3][7].value == "✗"
    assert ws[2][8].value in (None, "")
    assert ws[3][8].value in (None, "")


def test_maps_link_cell_is_clickable():
    ws = _load(build_sheet(_rows(), _exclusions()))["Places"]

    assert ws["G2"].hyperlink is not None
    assert ws["G2"].hyperlink.target == "https://www.google.com/maps/search/?api=1&query=Cafe+Alpha+Istanbul"
    assert ws["G3"].hyperlink.target == "https://www.google.com/maps/search/?api=1&query=Bistro+Beta+Istanbul"


def test_excluded_restaurants_sheet_lists_exclusions():
    wb = _load(build_sheet(_rows(), _exclusions()))
    ws = wb["Excluded restaurants"]

    assert [c.value for c in ws[1]] == ["Name", "Reason", "Detail"]
    assert [c.value for c in ws[2]] == ["Curry House", "diet", "Not vegetarian friendly"]
    assert ws.max_row == 2
