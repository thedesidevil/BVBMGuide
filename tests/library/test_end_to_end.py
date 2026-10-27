"""Seam test: real docx files -> build_plan -> CLI rebuild -> AIG v2 Library reader.

Only the model is stubbed (classifier client and ``extract_file``); everything
else is the real code path, so the shapes the prompt asks for must survive
clean, merge, shard naming and the v2 readers.
"""

import importlib.util
import json
from pathlib import Path

import pytest
from docx import Document
from typer.testing import CliRunner

import src.library.__main__ as cli
from src.aigv2.library import Library
from src.aigv2.writing.sections import _library_items
from src.library.ingest_core import runner as runner_mod
from src.library.ingest_core.extract import Extraction
from src.library.ingest_core.plan import build_plan

SCRIPT = Path(__file__).resolve().parents[2] / "scripts" / "check_library_for_v2.py"
FILLER = ("This is an obviously fake sample guide for Testville used only by the automated tests. "
          "It has enough text to pass the minimum length check of the extractor. ") * 3


def _docx(path: Path, day_table: bool = False) -> None:  # text differs per file (identical text is de-duplicated)
    doc = Document()
    doc.add_heading("Fake Guide to Rome", 1)
    doc.add_paragraph(FILLER + f" Marker {path.stem}.")
    for i in range(1, 4):  # day lines so the scanner treats it as an AIG
        doc.add_paragraph(f"Day {i}: Monday, {i} Jan - Fake Sightseeing in Rome")
    if day_table:
        t = doc.add_table(rows=3, cols=2)
        for i in range(3):
            t.cell(i, 0).text = f"Day {i + 1}"
            t.cell(i, 1).text = "Fake Trattoria lunch; Fake Ruins visit"
    doc.save(str(path))


def _extraction() -> Extraction:
    data = {
        "covered_cities": ["Rome"],
        "city_countries": {"Rome": "Italy"},
        "restaurants": [
            {"name": "Fake Trattoria", "city": "Rome", "cuisine_type": ["Roman"], "hours": "12-15",
             "meal_slots": ["lunch", "dinner"], "area": "Centro"},
            {"name": "Fake Cafe", "city": "Rome", "meal_slots": ["breakfast"]},
        ],
        "local_dishes": [{"name": "Fake Carbonara", "city": "Rome", "description": "Pasta"}],
        "connectivity_tips": [{"cities": ["Rome", "Florence"], "tip": "Public wifi is patchy near ruins"}],
        "transport_options": [
            {"city": "Rome", "mode": "Metro", "description": "Lines A, B", "recommended_pass": "Fake Roma Pass"},
            {"city": "Rome", "mode": "Train", "description": "Regional to Fakeville", "recommended_pass": ""},
            {"city": "Rome", "mode": "Train", "description": "High speed to Faketown", "recommended_pass": ""},
        ],
        "country_content": [
            {"country": "Italy", "currency": "EUR",
             "plugs": [{"type": "F", "voltage": "230V"}],
             "tipping": [{"tip": "Tipping is not expected in restaurants"}],
             "embassies": [{"name": "Fake Indian Embassy", "address": "1 Fake Via", "phone": "000"}],
             "connectivity_tips": [{"provider": "FakeSIM", "kind": "esim", "price_text": "EUR 10",
                                    "coverage_countries": ["Italy", "United States"]}],
             "phrases": [{"english": "Hello", "local": "Ciao", "category": "greeting"}],
             "emergency_contacts": [{"service": "Police", "number": "112"}]},
            {"country": "United States", "currency": "USD",
             "tipping": [{"tip": "Tip 18 to 20 percent"}],
             "connectivity_tips": [{"provider": "FakeSIM", "kind": "esim", "price_text": "EUR 10",
                                    "coverage_countries": ["Italy", "United States"]}]},
        ],
    }
    return Extraction(data=data, chunks=1, prompt_tokens=1, completion_tokens=1, usd=0.0)


class _Classifier:
    api_key = "k"

    def complete(self, prompt, **k):
        return '{"folder": "Italy"}'


class _Client:
    model = "stub-model"
    cost_usd = 0.0
    usage = {"prompt_tokens": 0, "completion_tokens": 0}


def test_docx_to_plan_to_rebuild_to_v2_reader(tmp_path, monkeypatch, capsys):
    inp, lib = tmp_path / "in", tmp_path / "lib"
    inp.mkdir()
    (lib / "Italy").mkdir(parents=True)
    _docx(inp / "a.docx", day_table=True)
    _docx(inp / "b.docx")
    _docx(inp / "c.docx")
    plan_dir = tmp_path / "plan"

    entries = build_plan(inp, lib, _Classifier(), plan_dir)
    assert len(entries) == 3 and {e.folder for e in entries} == {"Italy"}
    assert {e.fmt for e in entries} == {"old", "new"}

    monkeypatch.setattr("src.common.ai_provider.get_ai_client", lambda **k: _Client())
    monkeypatch.setattr(runner_mod, "extract_file", lambda c, p: _extraction())
    db = tmp_path / "new_db"
    res = CliRunner().invoke(cli.app, ["rebuild", "--plan", str(plan_dir), "--db", str(db),
                                       "--max-usd", "5", "--workers", "1"])
    assert res.exit_code == 0, res.output

    library = Library(db)
    rome = library.city("Rome")
    assert {r["name"] for r in rome["restaurants"]} == {"Fake Trattoria", "Fake Cafe"}
    assert next(r for r in rome["restaurants"] if r["name"] == "Fake Trattoria")["meal_slots"] == ["lunch", "dinner"]
    assert [t["tip"] for t in rome["connectivity_tips"]] == ["Public wifi is patchy near ruins"]
    assert len([t for t in rome["transport_options"] if t["mode"] == "Train"]) == 2

    usa = library.country("USA")
    assert usa["currency"] == "USD" and usa["tipping"][0]["tip"] == "Tip 18 to 20 percent"
    italy = library.country("Italy")
    assert italy["embassies"][0]["name"] == "Fake Indian Embassy"
    assert italy["currency"] == "EUR" and "city" not in italy["embassies"][0]
    assert library.country("United States") == {}  # only the canonical name is a shard
    assert sorted(p.stem for p in (db / "_country").glob("*.json")) == ["Italy", "USA"]

    cities, cdatas = [rome], [italy]
    getting = _library_items("getting_around", cities, cdatas)
    assert {t["mode"] for t in getting} == {"Metro", "Train"} and len(getting) == 3
    conn = _library_items("connectivity", cities, cdatas)
    assert {c.get("tip") or c.get("provider") for c in conn} == {"Public wifi is patchy near ruins", "FakeSIM"}
    assert _library_items("connectivity", [], [usa])[0]["provider"] == "FakeSIM"

    spec = importlib.util.spec_from_file_location("check_library_for_v2", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    capsys.readouterr()
    assert mod.main(["--db", str(db)]) == 0
    out = capsys.readouterr().out
    italy_row = next(l for l in out.splitlines() if l.startswith("Italy"))
    usa_row = next(l for l in out.splitlines() if l.startswith("USA"))
    assert italy_row.count("✓") >= 4 and "EUR" in italy_row
    assert "✓" in usa_row and "USD" in usa_row
