import io
import json
from datetime import time

import pytest
from docx import Document
from docx.oxml.ns import qn

from src.aigv2.library import Library
from datetime import date

from src.aigv2.models import DayPlan, Diet, HotelStay
from src.aigv2.planning.meals import MealPick
from src.aigv2.planning.slots import MealSlot
from src.aigv2.qc.reader import SECTION_TITLES, DocModel, read_docx
from src.aigv2.qc.rules import check_rules
from src.aigv2.render.blocks import add_link
from src.aigv2.render.document import render_guide
from src.aigv2.writing.days import RestaurantText
from tests.aigv2.test_render_document import HOTEL, S, chosen, data, day1, day2, day3, place, trip


@pytest.fixture
def library(tmp_path):
    shard = {"restaurants": [
        {"name": "Mikla", "cuisine_type": ["Turkish"], "vegetarian_friendly": False},
        {"name": "Hamdi", "cuisine_type": ["Turkish"], "vegetarian_friendly": False},
    ]}
    (tmp_path / "Istanbul.json").write_text(json.dumps(shard), encoding="utf-8")
    coverage = {"Turkey": ["Istanbul", "Cappadocia"], "Kerala": ["Kochi", "Munnar", "Alleppey"],
                "Bali": ["Bali", "Ubud", "Seminyak"]}
    (tmp_path / "_index.json").write_text(json.dumps({"_folder_coverage": coverage}), encoding="utf-8")
    return Library(tmp_path)


def render(d=None) -> bytes:
    return render_guide(d or data(), S)


def edit_text(docx: bytes, old: str, new: str) -> bytes:
    doc = Document(io.BytesIO(docx))
    hit = False
    for t in doc.element.body.iter(qn("w:t")):
        if t.text and old in t.text:
            t.text = t.text.replace(old, new)
            hit = True
    assert hit, old
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def rules_of(docx: bytes, library, t=None):
    return {f.rule for f in check_rules(read_docx(docx), t or trip(), library)}


def findings(docx: bytes, library, rule, t=None):
    return [f for f in check_rules(read_docx(docx), t or trip(), library) if f.rule == rule]


def with_intro(text: str, **kw):
    d1 = day1()
    d1.text.intro = text
    return data(days=[d1, day2(), day3()], **kw)


def google_roundtrip(docx: bytes, keep_headings: bool) -> bytes:
    """Stand-in for a Google Docs round-trip: custom styles become Normal, direct run formatting kept."""
    doc = Document(io.BytesIO(docx))
    by_id = {s.style_id: s for s in doc.styles if hasattr(s, "font")}
    for p in doc.element.body.iter(qn("w:p")):
        ppr = p.pPr
        ps = ppr.find(qn("w:pStyle")) if ppr is not None else None
        if ps is None:
            continue
        style = by_id.get(ps.get(qn("w:val")))
        if keep_headings and style is not None and style.name in ("Heading 1", "Heading 2", "Heading 3"):
            continue
        for r in p.iter(qn("w:r")):
            rpr = r.get_or_add_rPr()
            if style is not None and style.font.size and rpr.find(qn("w:sz")) is None:
                sz = rpr.makeelement(qn("w:sz"), {qn("w:val"): str(int(style.font.size.pt * 2))})
                rpr.append(sz)
            if style is not None and style.font.bold and rpr.find(qn("w:b")) is None:
                rpr.insert(0, rpr.makeelement(qn("w:b"), {}))
        ppr.remove(ps)
    return save(doc)


def snapshot(m: DocModel):
    return ([d.header for d in m.days],
            [[(s.heading, [v.name for v in s.venues]) for s in d.slots] for d in m.days],
            m.links, len(m.markers))


# ---------- reader ----------

def test_reader_extracts_days_slots_venues_facts_links_markers():
    raw = render()
    m = read_docx(raw)
    assert [d.header for d in m.days] == [
        "Day 1 – Thu, 29 Oct | Welcome to Istanbul", "Day 2 – Fri, 30 Oct | Sultans and Sunsets",
        "Day 3 – Sat, 31 Oct | Farewell, Istanbul"]
    d1 = m.days[0]
    assert [s.heading for s in d1.slots] == ["", "Afternoon", "Dinner near Pera Palace"]
    assert [v.name for v in d1.slots[0].venues] == ["Flight: Mumbai to Istanbul Airport", "Istanbul Airport → Pera Palace"]
    mikla = d1.slots[2].venues[0]
    assert mikla.name == "Mikla"
    assert mikla.url.startswith("https://www.google.com/maps/search/")
    assert mikla.facts["GETTING THERE"] == "Walk · 6 min"
    assert mikla.facts["PRICE"] == "€60 (~₹5,800)"
    assert "Rooftop views" in mikla.body
    assert any(m_.startswith("⚠ ADD RESTAURANT – Dinner, Day 1") for m_ in d1.slots[2].markers)
    assert [s.heading for s in m.days[1].slots][:2] == ["Morning", "Lunch — Included in your package"]
    assert ("official website", "https://topkapi.example") in m.links
    assert len(m.markers) == m.text.count("⚠")
    assert m.title == "Minarets and Moonlight"
    assert m.subtitle == "All Inclusive Guide – Istanbul"
    assert "Thank You" in m.sections and "Tailored Packing List" in m.sections
    assert [h for h in m.headings if h[0] == 1][-1] == (1, "Thank You")
    assert any(v.name == "Eczane Pera" for v in m.section_venues)


@pytest.mark.parametrize("keep_headings", [True, False], ids=["styles-lost", "headings-lost-too"])
def test_reader_survives_google_roundtrip(keep_headings):
    raw = render()
    original = read_docx(raw)
    trip_doc = read_docx(google_roundtrip(raw, keep_headings))
    assert snapshot(trip_doc) == snapshot(original)
    assert trip_doc.title == original.title
    assert [h for h in trip_doc.headings if h[0] == 1] == [h for h in original.headings if h[0] == 1]


@pytest.mark.parametrize("keep_headings", [True, False])
def test_roundtrip_doc_passes_rules(keep_headings, library):
    raw = google_roundtrip(render(), keep_headings)
    assert rules_of(raw, library) == {"markers_remaining"}


def test_reader_tolerates_legacy_doc(library):
    doc = Document()
    doc.add_paragraph("All Inclusive Guide- Tokyo")
    doc.add_paragraph("Day 1: Monday, 12 Jan – Arrival in Tokyo", style="Heading 1")
    doc.add_paragraph("Lunch Options:")
    t = doc.add_table(rows=2, cols=3)
    t.cell(0, 0).text = "Sushi Dai"
    doc.add_paragraph("Thank You", style="Heading 1")
    m = read_docx(save(doc))
    assert [d.header for d in m.days] == ["Day 1: Monday, 12 Jan – Arrival in Tokyo"]
    assert "day_header_format" in {f.rule for f in check_rules(m, trip(), library)}


def test_reader_finds_hand_added_restaurant_paragraph(library):
    doc = Document()
    doc.add_paragraph("Day 1 – Thu, 29 Oct | Welcome", style="Heading 1")
    doc.add_paragraph("Lunch near Galata Tower", style="Heading 2")
    p = doc.add_paragraph()
    add_link(p, "Karakoy Lokantasi", "https://www.google.com/maps/search/?api=1&query=Karakoy+Lokantasi")
    p.add_run(" – lovely meze")
    m = read_docx(save(doc))
    venues = m.days[0].slots[0].venues
    assert m.days[0].slots[0].heading == "Lunch near Galata Tower"
    assert [v.name for v in venues] == ["Karakoy Lokantasi"]
    assert venues[0].url.startswith("https://www.google.com/maps/")


# ---------- rules ----------

def test_rendered_guide_only_markers_remaining(library):
    raw = render()
    fs = check_rules(read_docx(raw), trip(), library)
    assert [f.rule for f in fs] == ["markers_remaining"]
    assert str(read_docx(raw).text.count("⚠")) in fs[0].message


def test_wrong_weekday(library):
    raw = edit_text(render(), "Day 1 – Thu, 29 Oct", "Day 1 – Fri, 29 Oct")
    fs = findings(raw, library, "day_weekday")
    assert fs and fs[0].day == 1 and fs[0].anchor_text == "Day 1 – Fri, 29 Oct | Welcome to Istanbul"
    assert not findings(raw, library, "day_dates")


def test_header_dates_differ_from_trip(library):
    raw = edit_text(render(), "Day 2 – Fri, 30 Oct", "Day 2 – Sat, 31 Oct")
    rules = rules_of(raw, library)
    assert "day_dates" in rules and "day_weekday" not in rules


def test_header_format(library):
    raw = edit_text(render(), "Day 1 – Thu, 29 Oct | Welcome to Istanbul", "Day 1: Thursday, 29 Oct – Welcome")
    fs = findings(raw, library, "day_header_format")
    assert fs and fs[0].day == 1


def test_section_order(library):
    m = read_docx(render())
    h = [x for x in m.headings if x != (1, "Thank You")]
    i = h.index((1, "Tailored Packing List"))
    m.headings = h[:i] + [(1, "Thank You")] + h[i:]
    assert "section_order" in {f.rule for f in check_rules(m, trip(), library)}
    m2 = read_docx(render())
    h2 = m2.headings
    a, b = h2.index((1, "Getting Around")), h2.index((1, "Tailored Packing List"))
    h2[a], h2[b] = h2[b], h2[a]
    assert "section_order" in {f.rule for f in check_rules(m2, trip(), library)}


def test_venue_without_maps_link(library):
    doc = Document(io.BytesIO(render()))
    for h in doc.element.body.iter(qn("w:hyperlink")):
        if "".join(t.text for t in h.iter(qn("w:t"))) == "Galata Tower":
            parent = h.getparent()
            for r in list(h):
                parent.insert(parent.index(h), r)
            parent.remove(h)
            break
    fs = findings(save(doc), library, "venue_link")
    assert [(f.anchor_text, f.day) for f in fs] == [("Galata Tower", 1)]


def test_venue_with_non_maps_link(library):
    d2 = day2()
    d2.places["Topkapi Palace"] = place("Topkapi Palace").model_copy(
        update={"maps_url": "https://www.tripadvisor.com/topkapi"})
    fs = findings(render(data(days=[day1(), d2, day3()])), library, "venue_link")
    assert [f.anchor_text for f in fs] == ["Topkapi Palace"]


def test_meal_with_too_few_restaurants(library):
    d2 = day2()
    d2.meals[0].restaurants = d2.meals[0].restaurants[:2]
    fs = findings(render(data(days=[day1(), d2, day3()])), library, "meal_count")
    assert len(fs) == 1 and fs[0].day == 2


def _lunch_day2(names, best_for=None):
    d2 = day2()
    lunch = MealSlot(meal="lunch", start=time(12, 30), anchors=[("Topkapi Palace", (41.0, 28.9), 3)])
    d2.result.meals = [lunch, d2.result.meals[1]]
    d2.meals = [MealPick(slot=lunch, day_index=2, restaurants=[chosen(n) for n in names])] + d2.meals
    if best_for:
        d2.text.restaurants.append(RestaurantText(name=names[0], ambience="Calm terrace", best_for=best_for))
    return d2


def test_repeated_restaurant_on_one_day(library):
    raw = render(data(days=[day1(), _lunch_day2(["Hamdi", "Kiva", "Pandeli"]), day3()]))
    rules = rules_of(raw, library)
    assert "repeat_day" in rules and "repeat_in_set" not in rules
    assert findings(raw, library, "repeat_day")[0].day == 2


def test_repeat_within_one_meal_set(library):
    d1 = day1()
    d1.meals[0].restaurants = [chosen("Mikla"), chosen("Mikla")]
    fs = findings(render(data(days=[d1, day2(), day3()])), library, "repeat_in_set")
    assert fs and fs[0].anchor_text == "Mikla" and fs[0].day == 1


def test_restaurant_more_than_twice_in_city(library):
    d2 = _lunch_day2(["Mikla", "Kiva", "Pandeli"])
    d3 = day3()
    lunch = MealSlot(meal="lunch", start=time(12, 0), anchors=[(HOTEL, (41.0, 28.9), 3)])
    d3.result.meals.append(lunch)
    d3.meals = [MealPick(slot=lunch, day_index=3, restaurants=[chosen("Mikla"), chosen("Asitane"), chosen("Lokanta")])]
    rules = rules_of(render(data(days=[day1(), d2, d3])), library)
    assert "repeat_city" in rules and "repeat_day" not in rules


def test_long_walk_fact_and_prose(library):
    d1 = day1()
    d1.meals[0].restaurants[0] = chosen("Mikla", minutes=25)
    fs = findings(render(data(days=[d1, day2(), day3()])), library, "walk_limit")
    assert any("25" in f.message and f.day == 1 for f in fs)
    raw = render(with_intro("The hotel is just a 15-minute walk from the tower."))
    assert findings(raw, library, "walk_limit")


def test_utm_link(library):
    d1 = day1()
    g = d1.places["Galata Tower"]
    d1.places["Galata Tower"] = g.model_copy(update={"maps_url": g.maps_url + "&utm_source=chatgpt.com"})
    assert "utm" in rules_of(render(data(days=[d1, day2(), day3()])), library)


def test_commentary_sentence_left_in(library):
    sentence = "Client is fussy about early starts so keep mornings light"
    t = trip(commentary_removed=[sentence])
    raw = render(with_intro(f"{sentence}.", trip=t))
    fs = findings(raw, library, "commentary", t)
    assert fs and fs[0].day == 1 and sentence in fs[0].anchor_text


@pytest.mark.parametrize("text", ["Dinner: [fill in later]", "Timings TBD", "Please mention timings here",
                                  "Lorem ipsum dolor", "Cuisine: Not specified", "Please see options below",
                                  "You can choose one of these cafes"])
def test_placeholders(text, library):
    fs = findings(render(with_intro(text)), library, "placeholder")
    assert fs and fs[0].anchor_text == text and fs[0].day == 1


@pytest.mark.parametrize("text", ["I'd keep dinner easy tonight after the long flight.",
                                  "Here are clear, step-by-step directions to your hotel.",
                                  "I would recommend the rooftop.", "Let me know if you want more options.",
                                  "Shown in your preferred format below."])
def test_assistant_chatter(text, library):
    assert findings(render(with_intro(text)), library, "assistant_chatter")


def test_house_text_here_are_passes(library):
    assert rules_of(render(with_intro("Here are few numbers you may need")), library) == {"markers_remaining"}


def test_title_must_name_destination(library):
    raw = edit_text(render(), "All Inclusive Guide – Istanbul", "All Inclusive Guide- Bali")
    fs = findings(raw, library, "title_destination")
    assert fs and "Bali" in fs[0].anchor_text


def test_closing_phrase_outside_thank_you(library):
    for text in ("We hope this vacation brings you lasting memories.", "Hope this trip brought you joy."):
        fs = findings(render(with_intro(text)), library, "closing_phrase_misplaced")
        assert fs and fs[0].day == 1


@pytest.mark.parametrize("text", ["Keep the client's passport handy.", "Ask the guest to carry cash.",
                                  "The clients' bags go first.", "Remind the client about the visa.",
                                  "The guest's room faces the sea."])
def test_third_person_client(text, library):
    assert findings(render(with_intro(text)), library, "third_person_client")


@pytest.mark.parametrize("text", ["The queen built her palace here.", "Locals greet him with tea.",
                                  "A clientele of artists gathers here."])
def test_third_person_ignores_plain_pronouns(text, library):
    assert not findings(render(with_intro(text)), library, "third_person_client")


def test_meal_label_mismatch_is_warning(library):
    raw = render(data(days=[day1(), _lunch_day2(["Kiva", "Pandeli", "Asitane"], best_for="Romantic dinner"), day3()]))
    fs = findings(raw, library, "meal_label_mismatch")
    assert [(f.severity, f.anchor_text, f.day) for f in fs] == [("warning", "Kiva", 2)]


def test_diet_only_checks_library_restaurants(library):
    t = trip(diet=Diet(restrictions=["vegetarian"]))
    fs = findings(render(data(trip=t)), library, "diet", t)
    assert {f.anchor_text for f in fs} == {"Mikla", "Hamdi"}
    assert all(f.severity == "error" for f in fs)


def test_markers_remaining_counts_by_type(library):
    m = read_docx(render())
    fs = [f for f in check_rules(m, trip(), library) if f.rule == "markers_remaining"]
    assert len(fs) == 1 and "ADD RESTAURANT" in fs[0].message and "VERIFY" in fs[0].message


def test_no_markers_no_marker_finding(library):
    doc = Document()
    doc.add_paragraph("Minarets", style="Title")
    doc.add_paragraph("All Inclusive Guide – Istanbul")
    for i, d in enumerate(trip().days, 1):
        doc.add_paragraph(f"Day {i} – {d.date:%a}, {d.date.day} {d.date:%b} | X", style="Heading 1")
    for title in SECTION_TITLES:
        doc.add_paragraph(title, style="Heading 1")
    assert check_rules(read_docx(save(doc)), trip(), library) == []


def test_anchors_are_verbatim_paragraphs(library):
    raw = render(with_intro("Timings TBD"))
    m = read_docx(raw)
    lines = set(m.text.splitlines())
    for f in check_rules(m, trip(), library):
        if f.anchor_text:
            assert f.anchor_text in lines, f


def test_walk_findings_per_venue_on_different_days(library):
    d1, d2 = day1(), day2()
    d1.meals[0].restaurants[0] = chosen("Mikla", minutes=25)
    d2.meals[0].restaurants[2] = chosen("Hamdi", minutes=25)
    fs = findings(render(data(days=[d1, d2, day3()])), library, "walk_limit")
    assert sorted((f.day, f.anchor_text) for f in fs) == [(1, "Mikla"), (2, "Hamdi")]


def test_walk_travel_line_attributed_to_next_card(library):
    d1 = day1()
    d1.legs["Galata Tower"] = d1.legs["Galata Tower"].model_copy(update={"walk_min": 9})
    raw = edit_text(render(data(days=[d1, day2(), day3()])), "Walk · 9 min", "Walk · 18 min")
    m = read_docx(raw)
    galata = next(v for d in m.days for s in d.slots for v in s.venues if v.name == "Galata Tower")
    assert galata.travel.startswith("Getting there: Walk · 18 min")
    fs = findings(raw, library, "walk_limit")
    assert [(f.anchor_text, f.day) for f in fs] == [("Galata Tower", 1)]


def _title_doc(title, subtitle):
    doc = Document()
    doc.add_paragraph(title, style="Title")
    doc.add_paragraph(subtitle)
    return read_docx(save(doc))


KERALA = dict(hotels=[HotelStay(city="Kochi", hotel="Brunton Boatyard", check_in=date(2026, 12, 1),
                                check_out=date(2026, 12, 3), source="voucher"),
                      HotelStay(city="Munnar", hotel="Blanket Hotel", check_in=date(2026, 12, 3),
                                check_out=date(2026, 12, 5), source="voucher")],
              days=[DayPlan(date=date(2026, 12, 1), title="Fort Kochi Walk", city="Kochi")],
              transport=[], trip_scope="domestic")


def _title_findings(title, subtitle, library, t):
    return [f for f in check_rules(_title_doc(title, subtitle), t, library) if f.rule == "title_destination"]


def test_region_title_passes(library):
    assert _title_findings("Kerala Escapes", "All Inclusive Guide – Kerala", library, trip(**KERALA)) == []


def test_unknown_region_title_is_only_a_warning(library, tmp_path_factory):
    bare = Library(tmp_path_factory.mktemp("bare"))
    fs = _title_findings("Kerala Escapes", "Your Southern Story", bare, trip(**KERALA))
    assert [f.severity for f in fs] == ["warning"]


def test_leaked_destination_in_title_is_error(library):
    fs = _title_findings("All Inclusive Guide - Bali", "", library, trip(**KERALA))
    assert [f.severity for f in fs] == ["error"] and "bali" in fs[0].message


def test_missing_mandatory_section(library):
    s = __import__("tests.aigv2.test_render_document", fromlist=["sections"]).sections()
    s["dishes"] = None
    fs = findings(render(data(sections=s)), library, "section_order")
    assert [f.message for f in fs] == ["Section 'Must-Try Local Dishes' is missing"]


def test_domestic_trip_omits_money_sim_health(library):
    s = __import__("tests.aigv2.test_render_document", fromlist=["sections"]).sections()
    for k in ("money", "connectivity", "health"):
        s[k] = None
    t = trip(trip_scope="domestic")
    assert not findings(render(data(sections=s, trip=t, trip_scope="domestic", domestic=True)), library,
                        "section_order", t)
    assert len(findings(render(data(sections=s)), library, "section_order")) == 3


def test_prose_paragraph_is_not_a_slot():
    doc = Document()
    doc.add_paragraph("Day 1 – Thu, 29 Oct | Welcome", style="Heading 1")
    doc.add_paragraph("Lunch near the river was lovely")
    doc.add_paragraph("Lunch near the river was lovely, and the views made it worth the walk from town.")
    p = doc.add_paragraph("Morning", style="List Bullet")
    doc.add_paragraph("Day 2 – Fri, 30 Oct | Old City was where we started the walk on our first visit years ago")
    m = read_docx(save(doc))
    assert m.days[0].slots == [] and [d.index for d in m.days] == [1]
    assert p.text in m.text


def test_hotel_name_not_in_vouchers(library):
    raw = edit_text(render(), "Istanbul Airport → Pera Palace", "Istanbul Airport → Grand Bazaar Inn")
    fs = findings(raw, library, "hotel_mismatch")
    assert [(f.anchor_text, f.day) for f in fs] == [("Istanbul Airport → Grand Bazaar Inn", 1)]


def test_flight_time_differs_from_ticket(library):
    raw = edit_text(render(), "04:00 · Mumbai", "05:00 · Mumbai")
    fs = findings(raw, library, "transport_time_mismatch")
    assert [(f.anchor_text, f.day) for f in fs] == [("Flight: Mumbai to Istanbul Airport", 1)]


def test_duplicate_and_out_of_order_days(library):
    raw = edit_text(render(), "Day 3 – Sat, 31 Oct", "Day 2 – Sat, 31 Oct")
    assert findings(raw, library, "day_sequence")
    m = read_docx(render())
    m.days = [m.days[1], m.days[0], m.days[2]]
    assert "day_sequence" in {f.rule for f in check_rules(m, trip(), library)}


def test_unknown_marker_types_counted():
    doc = Document()
    p = doc.add_paragraph("Hours 9–5 · ")
    p.add_run("⚠ CONFIRM WITH HOTEL")
    doc.add_paragraph("⚠ VERIFY: entry fee ⚠ PRICE UNKNOWN")
    assert read_docx(save(doc)).markers == ["⚠ CONFIRM WITH HOTEL", "⚠ VERIFY: entry fee", "⚠ PRICE UNKNOWN"]


def test_card_kinds_from_renderer_labels():
    m = read_docx(render())
    kinds = {v.name: v.kind for d in m.days for s in d.slots for v in s.venues}
    assert kinds["Flight: Mumbai to Istanbul Airport"] == "transport"
    assert kinds["Istanbul Airport → Pera Palace"] == "transfer"
    assert kinds["Mikla"] == kinds["Galata Tower"] == "venue"
