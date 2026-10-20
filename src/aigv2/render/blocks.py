from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_BREAK

from src.aigv2.models import Marker
from src.aigv2.render.theme import (
    BODY_FONT, CARD_BORDER, CARD_FILL, FROST, POWDER, ROYAL, SLATE, set_font,
)

_BM_ID = [100]


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def _style_run(run, *, size=None, bold=False, italic=False, color=None, caps=False, spacing=None):
    set_font(run, BODY_FONT)
    if size:
        run.font.size = Pt(size)
    run.font.bold = bold or None
    run.font.italic = italic or None
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if caps:
        run.font.small_caps = True
    if spacing:
        run._r.get_or_add_rPr().append(_el("w:spacing", val=int(spacing * 20)))
    _sort_rpr(run)


_RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "color",
              "spacing", "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd"]


def _sort_rpr(run):
    rpr = run._r.get_or_add_rPr()
    kids = list(rpr)
    for k in kids:
        rpr.remove(k)
    kids.sort(key=lambda k: _RPR_ORDER.index(k.tag.split("}")[1]) if k.tag.split("}")[1] in _RPR_ORDER else 99)
    for k in kids:
        rpr.append(k)


def _shade_cell(cell, fill):
    cell._tc.get_or_add_tcPr().append(_el("w:shd", val="clear", color="auto", fill=fill))


def _cell_margins(cell, top, left, bottom, right):
    mar = OxmlElement("w:tcMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        mar.append(_el(f"w:{side}", w=int(v * 20), type="dxa"))
    cell._tc.get_or_add_tcPr().append(mar)


def _table_borders(table, color=None, sz=4):
    pr = table._tbl.tblPr
    for old in pr.findall(qn("w:tblBorders")):
        pr.remove(old)
    b = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if color and name not in ("insideH", "insideV"):
            b.append(_el(f"w:{name}", val="single", sz=sz, space=0, color=color))
        else:
            b.append(_el(f"w:{name}", val="nil"))
    pr.append(b)


def _bottom_border(paragraph, color, sz=6):
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bdr.append(_el("w:bottom", val="single", sz=sz, space=1, color=color))
    ppr.append(bdr)


def add_link(paragraph, text: str, url: str, *, size=None, bold=False):
    rid = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    h.set(qn("w:history"), "1")
    _append_link_run(paragraph, h, text, size, bold)


def _append_link_run(paragraph, h, text, size, bold):
    run = paragraph.add_run(text)
    _style_run(run, size=size, bold=bold, color=ROYAL)
    run.font.underline = False
    h.append(run._r)
    paragraph._p.append(h)


def add_bookmark(paragraph, name: str):
    _BM_ID[0] += 1
    paragraph._p.insert(_first_content_index(paragraph), _el("w:bookmarkStart", id=_BM_ID[0], name=name))
    paragraph._p.append(_el("w:bookmarkEnd", id=_BM_ID[0]))


def _first_content_index(paragraph):
    ppr = paragraph._p.pPr
    return 0 if ppr is None else list(paragraph._p).index(ppr) + 1


def add_internal_link(paragraph, text: str, bookmark: str, *, size=None, bold=False):
    h = OxmlElement("w:hyperlink")
    h.set(qn("w:anchor"), bookmark)
    h.set(qn("w:history"), "1")
    _append_link_run(paragraph, h, text, size, bold)


def add_marker(paragraph, marker: Marker, *, size=None):
    run = paragraph.add_run(marker.render())
    _style_run(run, size=size, bold=True)
    rpr = run._r.get_or_add_rPr()
    rpr.append(_el("w:highlight", val="yellow"))
    _sort_rpr(run)
    return run


def add_eyebrow(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    _style_run(p.add_run(text), size=8.5, bold=True, color=ROYAL, caps=True, spacing=1.5)
    return p


def add_slot_heading(doc, text: str):
    p = doc.add_paragraph(text, style="Heading 2")
    p.paragraph_format.keep_with_next = True
    return p


def add_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    _bottom_border(p, POWDER)
    return p


def add_page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _add_value(paragraph, value, size=10):
    items = value if isinstance(value, list) else [value]
    for i, item in enumerate(items):
        if i:
            _style_run(paragraph.add_run(" · "), size=size, color=SLATE)
        if isinstance(item, Marker):
            add_marker(paragraph, item, size=size)
        else:
            _style_run(paragraph.add_run(str(item)), size=size, color=SLATE)


def _facts_table(cell, facts):
    t = cell.add_table(rows=(len(facts) + 1) // 2, cols=2)
    _table_borders(t)
    for i, (label, value) in enumerate(facts):
        c = t.cell(i // 2, i % 2)
        lp = c.paragraphs[0]
        lp.paragraph_format.space_after = Pt(0)
        lp.paragraph_format.keep_with_next = True
        _style_run(lp.add_run(label), size=7.5, bold=True, color=ROYAL, caps=True, spacing=1)
        vp = c.add_paragraph()
        vp.paragraph_format.space_after = Pt(4)
        _add_value(vp, value)
    return t


def add_card(doc, title, url, facts, body, pills):
    table = doc.add_table(rows=1, cols=1)
    _table_borders(table, CARD_BORDER)
    cell = table.cell(0, 0)
    _shade_cell(cell, CARD_FILL)
    _cell_margins(cell, 10, 12, 10, 12)
    tp = cell.paragraphs[0]
    tp.style = doc.styles["Heading 3"]
    tp.paragraph_format.keep_with_next = True
    if url:
        add_link(tp, title, url, size=15, bold=True)
        _heading_link_font(tp)
    else:
        tp.add_run(title)
    if facts:
        _facts_table(cell, facts)
    if body:
        bp = cell.add_paragraph()
        _style_run(bp.add_run(body), size=10, color=SLATE)
    if pills:
        pp = cell.add_paragraph()
        for i, pill in enumerate(pills):
            if i:
                pp.add_run("  ")
            r = pp.add_run(f" {pill} ")
            _style_run(r, size=8.5, color=ROYAL)
            r._r.get_or_add_rPr().append(_el("w:shd", val="clear", color="auto", fill=FROST))
            _sort_rpr(r)
    cell.paragraphs[-1].paragraph_format.space_after = Pt(0)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    return table


def _heading_link_font(paragraph):
    for r in paragraph._p.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        rf = rpr.find(qn("w:rFonts")) if rpr is not None else None
        if rf is not None:
            rpr.remove(rf)
        sz = rpr.find(qn("w:sz")) if rpr is not None else None
        if sz is not None:
            rpr.remove(sz)
            szcs = rpr.find(qn("w:szCs"))
            if szcs is not None:
                rpr.remove(szcs)
