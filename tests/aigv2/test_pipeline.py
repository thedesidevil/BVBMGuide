import copy
import io
import json
import re
from datetime import datetime, time
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from typer.testing import CliRunner

from src.aigv2 import pipeline
from src.aigv2.config import STEPS, Settings
from src.aigv2.library import Library
from src.aigv2.llm import StepLLM
from src.aigv2.models import Place, WeeklyHours
from src.aigv2.pipeline import Deps, UnansweredQuestions, build_run, cost_summary, regenerate, requalify, start_run
from src.library.ui.storage import LocalStorageBackend

FIX = Path(__file__).parent / "fixtures" / "nyc_trip"
VEG = ["Saravana Bhavan", "Dirt Candy", "Hangawi", "Bhavan Express", "Green Leaf", "Udupi Palace",
       "Leafy Table", "Garden Plate", "Sprout House", "Veg Corner"]


def _place(name: str, i: int = 0) -> Place:
    hours = WeeklyHours(periods={d: [(time(7), time(23))] for d in range(7)}, source="google")
    return Place(query=name, place_id=f"pid_{re.sub(r'\W', '_', name)}", name=name, address="1 Test St, New York, NY 10001, USA",
                 lat=40.75 + i * 0.0001, lng=-73.98, status="OPERATIONAL", hours=hours, website=None, phone="+1 212 000 0000",
                 maps_url=f"https://www.google.com/maps/search/?api=1&query={name.replace(' ', '+')}&query_place_id=pid",
                 fetched_at=datetime(2027, 1, 1), match_ok=True)


class FakePlaces:
    def __init__(self):
        self.calls = {"text_search": 0, "nearby": 0, "details": 0}

    def find(self, name, city, near=None, fresh=False):
        self.calls["text_search"] += 1
        return _place(name)

    def fresh_status(self, place):
        self.calls["details"] += 1
        return place

    def nearby(self, lat, lng, types, radius_m, max_results, city=None):
        self.calls["nearby"] += 1
        return [_place(f"{types[0]} {i}", i) for i in range(2)]


class FakeRoutes:
    def __init__(self):
        self.calls = {"matrix_elements": 0, "directions": 0}

    def durations(self, origins, destinations, mode):
        self.calls["matrix_elements"] += len(origins) * len(destinations)
        return {(i, j): 5 for i in range(len(origins)) for j in range(len(destinations))}

    def transit_steps(self, origin, destination):
        self.calls["directions"] += 1
        return []


class FakeFx:
    def rate(self, code, on=None):
        return 1.0 if code == "INR" else 85.0


def _rank(prompt: str) -> str:
    count = int(re.search(r"Choose exactly (\d+)", prompt).group(1))
    cands = json.loads(prompt.split("Candidates (JSON): ", 1)[1].split("\n", 1)[0])
    ordered = sorted(cands, key=lambda c: not c["indian"])
    return json.dumps({"picks": [c["name"] for c in ordered[:count]]})


def _day(prompt: str) -> str:
    facts = json.loads(prompt.split("\n", 1)[1])
    return json.dumps({
        "title": f"Exploring {facts['city']}", "intro": "We take it gently today.",
        "attractions": [{"name": a["name"], "description": "A favorite of ours."} for a in facts["attractions"]],
        "restaurants": [{"name": r["name"], "ambience": "Relaxed", "best_for": "Families"} for r in facts["restaurants"]],
        "free_time_idea": None})


REPLIES = {
    "extract": lambda p: (FIX / "expected_trip_plan.json").read_text(encoding="utf-8"),
    "rank": _rank,
    "write_days": _day,
    "write_sections": lambda p: json.dumps({"items": [{"heading": "Good to know", "body": "Plan ahead and enjoy."}]}),
    "title": lambda p: json.dumps({"title": "Manhattan Lights Ahead"}),
    "qc_prose": lambda p: json.dumps({"issues": []}),
    "qc_sense": lambda p: json.dumps({"issues": []}),
}


class FakeLLM(StepLLM):
    def __init__(self, settings):
        super().__init__(settings, client_factory=lambda m: None)
        self.calls: list[tuple[str, str]] = []

    def _attempt(self, step, model, system, prompt, schema, validate, max_tokens, fallback):
        self.calls.append((step, prompt))
        self._log.append({"step": step, "model": model, "prompt_tokens": 100, "completion_tokens": 20,
                          "usd": 0.001, "fallback": fallback})
        obj = schema.model_validate_json(REPLIES[step](prompt))
        problems = validate(obj) if validate else []
        return (None, problems) if problems else (obj, [])


def _library(root: Path) -> Library:
    root.mkdir()
    rests = [{"name": n, "cuisine_type": ["South Indian"] if "Bhavan" in n or "Udupi" in n else ["Vegetarian"],
              "vegetarian_friendly": True, "highlights": "Pure vegetarian kitchen", "price_range": "$20-30"}
             for n in VEG]
    rests.append({"name": "Prime Steakhouse", "cuisine_type": ["Steak"], "vegetarian_friendly": False,
                  "price_range": "$80"})
    shard = {"restaurants": rests, "attractions": [{"name": "Statue of Liberty ferry", "recommended_duration": "3 hours"}],
             "safety_tips": [{"tip": "Keep valuables close."}]}
    (root / "New York.json").write_text(json.dumps(shard), encoding="utf-8")
    (root / "_index.json").write_text(json.dumps({"_folder_coverage": {"USA": ["New York"]}}), encoding="utf-8")
    (root / "_country").mkdir()
    (root / "_country" / "USA.json").write_text(json.dumps({"currency": "USD"}), encoding="utf-8")
    return Library(root)


@pytest.fixture
def deps(tmp_path):
    models = {s: "gpt-6-luna" for s in STEPS} | {"qc_sense": "gpt-6.1-sol"}
    settings = Settings(models=models, data_root=tmp_path / "data")
    return Deps(settings=settings, llm=FakeLLM(settings), library=_library(tmp_path / "lib"), places=FakePlaces(),
                routes=FakeRoutes(), fx=FakeFx(), backend=LocalStorageBackend(tmp_path / "data"))


def answers_for(questions) -> dict:
    out = {}
    for q in questions:
        if q.kind == "vague_day":
            out[q.id] = "Museum of Modern Art"
        else:
            out[q.id] = q.options[0] if q.options else "yes"
    return out


def started(deps):
    store, questions = start_run(None, sorted(FIX.glob("*.txt")), deps)
    return store, questions, answers_for(questions)


def built(deps):
    store, questions, answers = started(deps)
    return store, answers, build_run(store, answers, deps)


def day_headers(docx: bytes) -> list[str]:
    doc = Document(io.BytesIO(docx))
    return [p.text for p in doc.paragraphs if p.style.name == "Heading 1" and p.text.startswith("Day ")]


def steps(llm, name):
    return [p for s, p in llm.calls if s == name]


def test_end_to_end_fakes(deps):
    store, questions, answers = started(deps)
    assert questions and {q.kind for q in questions} >= {"open_choice", "vague_day"}
    assert store.load_json("trip_plan.json")["client_names"] == ["Test Client"]
    assert store.load_json("questions.json")
    result = build_run(store, answers, deps)
    for name in ("guide.docx", "qc_report.md", "qc_report.json", "verification.xlsx", "cost.json", "plan.json",
                 "days/day_01.json", "sections/packing.json"):
        assert store.has(name), name
    headers = day_headers(store.load_bytes("guide.docx"))
    assert len(headers) == 4
    assert [int(re.match(r"Day (\d+) – ", h).group(1)) for h in headers] == [1, 2, 3, 4]
    assert headers[0].startswith("Day 1 – Fri, 12 Mar | ")
    assert result.docx_name == "All Inclusive Guide - Test Client - New York.docx"
    assert result.cost_usd > 0 and result.qc.marker_count >= 0
    assert "Museum of Modern Art" in "\n".join(steps(deps.llm, "write_days"))
    assert "Prime Steakhouse" not in "\n".join(p for _, p in deps.llm.calls if _ in ("rank", "write_days"))


def test_build_refuses_unanswered(deps):
    store, questions, answers = started(deps)
    with pytest.raises(UnansweredQuestions) as exc:
        build_run(store, {}, deps)
    assert {q.id for q in exc.value.questions} == {q.id for q in questions}
    partial = dict(answers)
    partial.pop(next(iter(partial)))
    with pytest.raises(UnansweredQuestions):
        build_run(store, partial, deps)
    assert not store.has("plan.json")


def test_checkpoint_skips_stage(deps):
    store, answers, first = built(deps)
    before = len(deps.llm.calls)
    second = build_run(store, answers, deps)
    assert len(deps.llm.calls) == before
    assert second.docx_name == first.docx_name and second.qc.passed == first.qc.passed
    assert second.cost_usd == 0


def test_regenerate_one_day(deps):
    store, answers, _ = built(deps)
    deps.llm.calls.clear()
    old = store.load_bytes("guide.docx")
    result = regenerate(store, [2], [], deps)
    writes = steps(deps.llm, "write_days")
    assert len(writes) == 1 and '"date": "2027-03-13"' in writes[0]
    assert not steps(deps.llm, "write_sections") and not steps(deps.llm, "title")
    assert steps(deps.llm, "qc_sense")
    assert store.load_bytes("guide.v1.docx") == old
    assert store.has("guide.docx") and store.has("qc_report.v1.md")
    assert len(day_headers(store.load_bytes("guide.docx"))) == 4
    assert result.docx_name.endswith(".docx")
    regenerate(store, [], ["packing"], deps)
    assert store.has("guide.v2.docx")
    assert len(steps(deps.llm, "write_sections")) == 1


def test_regenerate_keeps_repeat_limits(deps):
    store, answers, _ = built(deps)
    regenerate(store, [2], [], deps)
    plan = store.load_json("plan.json")
    uses: dict = {}
    for pick in plan["picks"]:
        names = [c["lib"]["name"] for c in pick["restaurants"]]
        assert len(names) == len(set(names))
        for n in names:
            uses[n] = uses.get(n, 0) + 1
    assert uses and max(uses.values()) <= 2


def test_regenerate_rejects_unknown(deps):
    store, answers, _ = built(deps)
    with pytest.raises(ValueError):
        regenerate(store, [9], [], deps)
    with pytest.raises(ValueError):
        regenerate(store, [], ["nonsense"], deps)


def _strip_markers(doc) -> None:
    for p in list(doc.element.body.iter(qn("w:p"))):
        text = "".join(t.text or "" for t in p.iter(qn("w:t")))
        if "⚠" not in text:
            continue
        if p.getparent() is doc.element.body:
            p.getparent().remove(p)
        else:
            for t in p.iter(qn("w:t")):
                if t.text and "⚠" in t.text:
                    t.text = "checked"


def _add_restaurant(doc, name: str) -> None:
    for tbl in doc.element.body.iter(qn("w:tbl")):
        text = "".join(t.text or "" for t in tbl.iter(qn("w:t")))
        hit = next((n for n in VEG if n in text), None)
        if hit and "CUISINE" in text.upper():
            new = copy.deepcopy(tbl)
            for t in new.iter(qn("w:t")):
                if t.text and hit in t.text:
                    t.text = t.text.replace(hit, name)
            tbl.addnext(new)
            return
    raise AssertionError("no restaurant card found")


def test_requalify_edited_doc(deps, tmp_path):
    store, answers, _ = built(deps)
    doc = Document(io.BytesIO(store.load_bytes("guide.docx")))
    _strip_markers(doc)
    _add_restaurant(doc, "Hand Added Bistro")
    path = tmp_path / "edited.docx"
    doc.save(str(path))
    out, result, report = requalify(store, path, deps)
    assert out[:2] == b"PK" and "# QC report" in report
    mine = [f for f in result.findings if f.anchor_text == "Hand Added Bistro" or "Hand Added Bistro" in f.message]
    assert not [f for f in mine if f.rule == "diet" or "library" in f.message.casefold()]
    assert "Hand Added Bistro" in [r.venue for r in result.fact_rows]
    assert not any(f.rule == "diet" for f in result.findings)


def test_cost_log_totals(deps):
    store, answers, result = built(deps)
    cost = store.load_json("cost.json")
    assert cost["ai"]["total_usd"] + cost["google"]["total_usd"] == pytest.approx(cost["total_usd"])
    assert result.cost_usd == pytest.approx(cost["total_usd"])
    by_step, by_api = cost["ai"]["by_step"], cost["google"]["by_api"]
    assert {"write_days", "write_sections", "title", "qc_prose", "qc_sense"} <= set(by_step)
    assert by_step["qc_sense"]["model"] == "gpt-6.1-sol" and by_step["write_days"]["prompt_tokens"] > 0
    assert sum(s["usd"] for s in by_step.values()) == pytest.approx(cost["ai"]["total_usd"])
    unit = deps.settings.google_usd_per_call
    assert set(unit) == {"text_search", "nearby", "details", "route_element"}
    for api in ("text_search", "nearby", "details", "route_element"):
        assert by_api[api]["calls"] > 0
        assert by_api[api]["usd"] == pytest.approx(by_api[api]["calls"] * unit[api])
    assert by_api["route_element"]["calls"] == deps.routes.calls["matrix_elements"] + deps.routes.calls["directions"]
    assert sum(a["usd"] for a in by_api.values()) == pytest.approx(cost["google"]["total_usd"])
    text = cost_summary(cost)
    assert text.startswith(f"Cost this run: ${cost['total_usd']:.2f}")
    assert "write_days $" in text and "text_search " in text and " calls $" in text


def test_cli_prints_cost_split(deps, tmp_path, monkeypatch):
    from src.aigv2.__main__ import app
    store, questions, answers = started(deps)
    ans = tmp_path / "answers.json"
    ans.write_text(json.dumps(answers), encoding="utf-8")
    monkeypatch.setattr(pipeline, "default_deps", lambda: deps)
    out = CliRunner().invoke(app, ["build", store.run_id, "--answers", str(ans)])
    assert out.exit_code == 0, out.output
    lines = out.output.splitlines()
    assert any(l.startswith("Cost this run: $") for l in lines)
    assert any(l.strip().startswith("AI ") and "write_days" in l for l in lines)
    assert any(l.strip().startswith("Google ") and "text_search" in l for l in lines)
    assert "PASS" in out.output or "FAIL" in out.output
    assert "guide.docx" in out.output and "Markers:" in out.output


def test_cli_start_and_qc(deps, tmp_path, monkeypatch):
    from src.aigv2.__main__ import app
    monkeypatch.setattr(pipeline, "default_deps", lambda: deps)
    args = ["start", *[str(p) for p in sorted(FIX.glob("*.txt"))]]
    out = CliRunner().invoke(app, args)
    assert out.exit_code == 0, out.output
    assert "Run id:" in out.output and re.search(r"^ *1\. ", out.output, re.M) and "Cost this run:" in out.output
    store, answers, _ = built(deps)
    edited = tmp_path / "edited.docx"
    edited.write_bytes(store.load_bytes("guide.docx"))
    out = CliRunner().invoke(app, ["qc", str(edited), "--run", store.run_id])
    assert out.exit_code == 0, out.output
    assert "Cost this run:" in out.output and ("PASS" in out.output or "FAIL" in out.output)
    assert (tmp_path / "edited.qc.docx").exists() and (tmp_path / "edited.qc_report.md").exists()
    assert json.loads((tmp_path / "edited.cost.json").read_text())["total_usd"] >= 0


# ---------- fix round 1 ----------

from datetime import date, timedelta  # noqa: E402

from src.aigv2.models import DayPlan, Diet, HotelStay, PlannedItem, TransportLeg, Traveller, TripPlan  # noqa: E402
from src.aigv2.pipeline import pkey  # noqa: E402
from src.aigv2.preflight import run_checklist  # noqa: E402
from src.aigv2.intake.template import IntakeFields  # noqa: E402
from src.aigv2.store import RunStore  # noqa: E402
from src.aigv2.writing.sections import write_section  # noqa: E402

D0 = date(2027, 5, 10)
COUNTRY = {"Prague": "Czechia", "Vienna": "Austria"}


class CityPlaces(FakePlaces):
    def __init__(self):
        super().__init__()
        self.lookups = []
        self.closed: set[tuple[str, str]] = set()

    def find(self, name, city, near=None, fresh=False):
        self.calls["text_search"] += 1
        self.lookups.append((name, city))
        p = _place(name)
        p.lat = 50.0 if city == "Prague" else 48.2 if city == "Vienna" else 40.75
        p.address = f"1 Test St, {city}, {COUNTRY.get(city, 'USA')}"
        p.place_id = f"pid_{name}_{city}".replace(" ", "_")
        p.lng = -73.98 + len(name) * 0.001
        if (name, city) in self.closed:
            p.hours = WeeklyHours(periods={6: [(time(9), time(17))]}, source="google")
        return p


def two_city_trip() -> TripPlan:
    d = [D0 + timedelta(days=i) for i in range(4)]
    return TripPlan(
        client_names=["Test Client"], travellers=Traveller(adults=2), diet=Diet(), travel_modes=["taxi"],
        hotels=[HotelStay(city="Prague", hotel="Grand Hotel", check_in=d[0], check_out=d[2], source="voucher"),
                HotelStay(city="Vienna", hotel="Grand Hotel", check_in=d[2], check_out=d[3], source="voucher")],
        transport=[
            TransportLeg(kind="flight", from_place="Mumbai", to_place="Airport", depart=datetime(2027, 5, 10, 2, 0),
                         arrive=datetime(2027, 5, 10, 9, 0), source="ticket"),
            TransportLeg(kind="flight", from_place="Airport", to_place="Airport", depart=datetime(2027, 5, 12, 10, 0),
                         arrive=datetime(2027, 5, 12, 11, 0), source="ticket"),
            TransportLeg(kind="flight", from_place="Airport", to_place="Mumbai", depart=datetime(2027, 5, 13, 18, 0),
                         arrive=datetime(2027, 5, 14, 6, 0), source="ticket")],
        bookings=[],
        days=[DayPlan(date=d[0], title="Arrive", city="Prague"),
              DayPlan(date=d[1], title="Old Town", city="Prague",
                      items=[PlannedItem(name="Old Town", kind="attraction", city="Prague", slot_hint="morning"),
                             PlannedItem(name="Castle", kind="attraction", city="Prague", slot_hint="afternoon")]),
              DayPlan(date=d[2], title="To Vienna", city="Vienna",
                      items=[PlannedItem(name="Old Town", kind="attraction", city="Vienna", slot_hint="afternoon")]),
              DayPlan(date=d[3], title="Home", city="Vienna")],
        open_choices=[], vague_days=[], commentary_removed=[], trip_scope="international")


def seeded(deps, trip: TripPlan) -> RunStore:
    store = RunStore.create(deps.backend)
    store.save_json("trip_plan.json", trip)
    store.save_json("questions.json", {"questions": []})
    return store


def test_places_keyed_by_city(deps):
    deps.places = CityPlaces()
    store = seeded(deps, two_city_trip())
    build_run(store, {}, deps)
    state = pipeline._load_plan(store)
    d2, d3, d4 = state.days[1], state.days[2], state.days[3]
    assert d2.places["Old Town"].lat == 50.0 and d3.places["Old Town"].lat == 48.2
    assert d3.places["Grand Hotel"].lat == 48.2 and d2.places["Grand Hotel"].lat == 50.0
    assert d4.places["Airport"].lat == 48.2 and state.days[0].places["Airport"].lat == 50.0
    assert ("Airport", "Vienna") in deps.places.lookups and ("Airport", "Prague") in deps.places.lookups
    assert state.countries == {"Prague": "Czechia", "Vienna": "Austria"}
    assert {i.city for i in state.important} == {"Prague", "Vienna"}


class RecordingRoutes(FakeRoutes):
    def __init__(self):
        super().__init__()
        self.pairs = []

    def durations(self, origins, destinations, mode):
        self.pairs.append((list(origins), list(destinations)))
        return super().durations(origins, destinations, mode)


def test_arrival_departure_and_legs(deps):
    deps.places, deps.routes = CityPlaces(), RecordingRoutes()
    store = seeded(deps, two_city_trip())
    build_run(store, {}, deps)
    state = pipeline._load_plan(store)
    assert [(d.arrival, d.departure) for d in state.days] == [(0, None), (None, None), (1, 1), (None, 2)]
    assert set(state.days[0].legs) == {"arrival"}
    assert set(state.days[1].legs) == {"Old Town", "Castle"}
    assert state.days[1].legs["Castle"].from_name == "Old Town"
    assert set(state.days[2].legs) == {"arrival", "departure", "Old Town"}
    assert state.days[2].legs["departure"].from_name == "Grand Hotel"
    assert state.days[2].legs["arrival"].to_name == "Grand Hotel"
    view = state.days[2].places
    assert view[pkey("Grand Hotel", "Prague")].lat == 50.0 and view[pkey("Grand Hotel", "Vienna")].lat == 48.2
    assert view[pkey("Airport", "Prague")].lat == 50.0 and view[pkey("Airport", "Vienna")].lat == 48.2
    assert view["Grand Hotel"].lat == 48.2
    prague_hotel, prague_port = (50.0, view[pkey("Grand Hotel", "Prague")].lng), (50.0, view[pkey("Airport", "Prague")].lng)
    vienna_port, vienna_hotel = (48.2, view[pkey("Airport", "Vienna")].lng), (48.2, view[pkey("Grand Hotel", "Vienna")].lng)
    assert ([prague_hotel], [prague_port]) in deps.routes.pairs
    assert ([vienna_port], [vienna_hotel]) in deps.routes.pairs
    assert set(state.days[3].legs) == {"departure"}
    rds = pipeline._render_days(store, state, deps)
    assert rds[2].arrival is rds[2].departure and rds[2].prev_city == "Prague"
    assert rds[0].departure is None and rds[3].arrival is None
    assert rds[2].utc_offset_hours == 2.0


def test_early_arrival_and_after_midnight_departure():
    d = [D0 + timedelta(days=i) for i in range(3)]
    trip = two_city_trip().model_copy(update={
        "hotels": [HotelStay(city="Prague", hotel="H", check_in=d[0], check_out=d[2], source="voucher")],
        "transport": [
            TransportLeg(kind="flight", from_place="X", to_place="Airport", arrive=datetime(2027, 5, 11, 2, 30),
                         source="t"),
            TransportLeg(kind="flight", from_place="Airport", to_place="X", depart=datetime(2027, 5, 13, 1, 30),
                         source="t")]})
    assert pipeline._arrival_idx(trip, d[1]) == 0
    assert pipeline._arrival_idx(trip, d[0]) is None
    assert pipeline._departure_idx(trip, d[2], {d[0], d[1], d[2]}) == 1
    assert pipeline._departure_idx(trip, d[2], {d[0], d[1], d[2], d[2] + timedelta(days=1)}) is None
    late = trip.model_copy(update={"transport": [trip.transport[1].model_copy(
        update={"depart": datetime(2027, 5, 13, 6, 0)})]})
    assert pipeline._departure_idx(late, d[2], {d[2]}) is None


def test_utc_offset_and_country():
    assert pipeline._utc_offset("Austria", date(2027, 1, 15)) == 1.0
    assert pipeline._utc_offset("Austria", date(2027, 7, 15)) == 2.0
    assert pipeline._utc_offset("India", date(2027, 7, 15)) == 5.5
    assert pipeline._utc_offset("USA", date(2027, 7, 15)) is None
    assert pipeline._utc_offset("", date(2027, 7, 15)) is None
    p = _place("Hotel")
    p.address = "Kärntner Ring 1, 1010 Wien, Austria"
    assert pipeline._country_of(p) == "Austria"
    p.address = "10 Bayfront Ave, Singapore 018956"
    assert pipeline._country_of(p) == "Singapore"
    p.match_ok = False
    assert pipeline._country_of(p) == "" and pipeline._country_of(None) == ""


def test_country_falls_back_to_library(deps, tmp_path):
    assert pipeline._library_country(deps.library, "New York") == "USA"
    assert pipeline._library_country(deps.library, "Atlantis") == ""

    class NoCountry(FakePlaces):
        def find(self, name, city, near=None, fresh=False):
            p = super().find(name, city, near, fresh)
            p.address = ""
            return p
    deps.places = NoCountry()
    store, answers, _ = built(deps)
    assert pipeline._load_plan(store).countries == {"New York": "USA"}


def test_regenerate_swaps_exclusions(deps):
    store, answers, _ = built(deps)
    plan = store.load_json("plan.json")
    plan["exclusions"] += [{"name": "S2", "reason": "closed", "detail": "Lunch, Day 2, near X: sentinel"},
                           {"name": "S3", "reason": "closed", "detail": "Lunch, Day 3, near X: sentinel"}]
    store.save_json("plan.json", plan)
    regenerate(store, [2], [], deps)
    names = [e["name"] for e in store.load_json("plan.json")["exclusions"]]
    assert "S2" not in names and "S3" in names
    assert any(", Day 2, " in e["detail"] for e in store.load_json("plan.json")["exclusions"])


def test_changed_answers_refused(deps):
    store, answers, _ = built(deps)
    changed = dict(answers)
    changed["open_choice:0"] = "Top of the Rock sunset visit"
    with pytest.raises(ValueError, match="start a new run"):
        build_run(store, changed, deps)


def test_failed_build_saves_partial_and_cost(deps, tmp_path, monkeypatch):
    store, questions, answers = started(deps)

    def boom(prompt):
        raise RuntimeError("title service down")
    monkeypatch.setitem(REPLIES, "title", boom)
    with pytest.raises(RuntimeError) as exc:
        build_run(store, answers, deps)
    assert exc.value.run_id == store.run_id and exc.value.run_cost["ai"]["total_usd"] > 0
    assert store.has("days/day_01.json") and store.has("days/day_04.json") and store.has("sections/safety.json")
    assert not store.has("sections/title.json") and not store.has("guide.docx")
    cost = store.load_json("cost.json")
    assert cost == exc.value.run_cost and cost["google"]["by_api"]["text_search"]["calls"] > 0
    assert store.load_json("cost_history.json")["runs"][-1]["command"] == "build (failed)"
    ans = tmp_path / "answers.json"
    ans.write_text(json.dumps(answers), encoding="utf-8")
    from src.aigv2.__main__ import app
    monkeypatch.setattr(pipeline, "default_deps", lambda: deps)
    out = CliRunner().invoke(app, ["build", store.run_id, "--answers", str(ans)])
    assert out.exit_code == 1 and "FAILED: RuntimeError: title service down" in out.output
    assert "Cost this run: $" in out.output and "  Google " in out.output
    monkeypatch.setitem(REPLIES, "title", lambda p: json.dumps({"title": "Manhattan Lights Ahead"}))
    before = len(steps(deps.llm, "write_days"))
    build_run(store, answers, deps)
    assert len(steps(deps.llm, "write_days")) == before and store.has("guide.docx")


def test_requalify_requires_run(deps, tmp_path, monkeypatch):
    store, answers, _ = built(deps)
    edited = tmp_path / "edited.docx"
    edited.write_bytes(store.load_bytes("guide.docx"))
    with pytest.raises(ValueError, match="--run"):
        requalify(None, edited, deps)
    from src.aigv2.__main__ import app
    monkeypatch.setattr(pipeline, "default_deps", lambda: deps)
    out = CliRunner().invoke(app, ["qc", str(edited)])
    assert out.exit_code != 0 and "--run" in out.output


def test_domestic_guide_has_health_check_marker(deps, monkeypatch):
    plan = json.loads((FIX / "expected_trip_plan.json").read_text(encoding="utf-8"))
    plan["trip_scope"] = "domestic"
    monkeypatch.setitem(REPLIES, "extract", lambda p: json.dumps(plan))
    store, answers, result = built(deps)
    text = "\n".join(p.text for p in Document(io.BytesIO(store.load_bytes("guide.docx"))).paragraphs)
    assert "⚠ CHECK: Altitude / remote area / sea travel not assessed" in text
    assert result.qc.marker_kinds.get("CHECK", 0) >= 1


def test_health_section_for_high_altitude(deps):
    plan = TripPlan.model_validate_json((FIX / "expected_trip_plan.json").read_text(encoding="utf-8"))
    domestic = plan.model_copy(update={"trip_scope": "domestic"})
    city = [{"city": "Leh", "country": "India"}]
    assert write_section("health", city, {}, domestic, deps.llm) is None
    assert write_section("health", city, {}, domestic, deps.llm, high_altitude=True) is not None
    assert write_section("money", city, {}, domestic, deps.llm, high_altitude=True) is None


def test_checklist_hours_checked_per_city(deps, monkeypatch):
    trip = two_city_trip()
    places = CityPlaces()
    places.closed = {("Old Town", "Prague")}
    resolved = pipeline.resolve_places(trip, places, {})
    qs = run_checklist(trip, IntakeFields(), [], deps.library, resolved, deps.settings)
    hours = [q for q in qs if q.kind == "hours_conflict"]
    assert [(q.id, q.day) for q in hours] == [(f"hours_conflict:{D0 + timedelta(days=1)}:Old Town", D0 + timedelta(days=1))]
    places.closed = {("Old Town", "Vienna")}
    monkeypatch.setitem(REPLIES, "extract", lambda p: trip.model_dump_json())
    deps.places = places
    _, qs = start_run(None, sorted(FIX.glob("*.txt")), deps)
    assert [q.day for q in qs if q.kind == "hours_conflict"] == [D0 + timedelta(days=2)]


# ---------- final review fixes ----------

from src.aigv2.qc.reader import read_docx  # noqa: E402
from src.aigv2.qc.rules import check_rules  # noqa: E402

M = [date(2027, 3, 10) + timedelta(days=i) for i in range(4)]


def day_trip() -> TripPlan:
    return TripPlan(
        client_names=["Test Client"], travellers=Traveller(adults=2), diet=Diet(), travel_modes=["taxi"],
        hotels=[HotelStay(city="Florence", hotel="Hotel Arno", check_in=M[0], check_out=M[3], source="voucher")],
        transport=[
            TransportLeg(kind="flight", from_place="Mumbai", to_place="Florence Airport",
                         depart=datetime(2027, 3, 10, 2, 0), arrive=datetime(2027, 3, 10, 10, 0), source="ticket"),
            TransportLeg(kind="train", from_place="Firenze SMN", to_place="Pisa Centrale",
                         depart=datetime(2027, 3, 11, 9, 0), arrive=datetime(2027, 3, 11, 10, 0), source="ticket"),
            TransportLeg(kind="train", from_place="Pisa Centrale", to_place="Firenze SMN",
                         depart=datetime(2027, 3, 11, 18, 0), arrive=datetime(2027, 3, 11, 19, 0), source="ticket"),
            TransportLeg(kind="flight", from_place="Florence Airport", to_place="Mumbai",
                         depart=datetime(2027, 3, 13, 15, 0), arrive=datetime(2027, 3, 13, 23, 0), source="ticket")],
        bookings=[],
        days=[DayPlan(date=M[0], title="Arrive", city="Florence"),
              DayPlan(date=M[1], title="Day trip to Pisa", city="Florence",
                      items=[PlannedItem(name="Leaning Tower", kind="attraction", city="Pisa", time="11:00")]),
              DayPlan(date=M[2], title="Uffizi", city="Florence"),
              DayPlan(date=M[3], title="Home", city="Florence")],
        open_choices=[], vague_days=[], commentary_removed=[], trip_scope="international")


def day_text(docx: bytes, n: int) -> str:
    m = read_docx(docx)
    return "\n".join(p.text for p in m.paragraphs if p.day == n)


def test_same_hotel_day_trip_is_not_a_departure(deps):
    trip = day_trip()
    days = {d.date for d in trip.days}
    assert pipeline._departure_idx(trip, M[1], days) is None
    assert pipeline._departure_idx(trip, M[3], days) == 3
    store = seeded(deps, trip)
    result = build_run(store, {}, deps)
    state = pipeline._load_plan(store)
    d2 = state.days[1]
    assert (d2.arrival, d2.departure) == (None, None) and d2.other_legs == [1, 2]
    assert {m.meal for m in d2.result.meals} >= {"lunch", "dinner"}
    assert "arrival" not in d2.legs and "departure" not in d2.legs
    docx = store.load_bytes("guide.docx")
    text = day_text(docx, 2)
    assert "Reach the station by" not in text and "Overnight in Florence" in text
    assert "Train: Firenze SMN to Pisa Centrale" in text and "Train: Pisa Centrale to Firenze SMN" in text
    assert "⚠ VERIFY: day-trip transport" in text
    assert "Lunch near" in text and "Dinner near" in text
    assert not [f for f in result.qc.findings if f.rule == "transport_missing"]


def test_qc_flags_transport_leg_not_in_guide(deps):
    trip = day_trip()
    store = seeded(deps, trip)
    build_run(store, {}, deps)
    docx = store.load_bytes("guide.docx")
    doc = Document(io.BytesIO(docx))
    for tbl in list(doc.element.body.iter(qn("w:tbl"))):
        if "Train: Pisa Centrale to Firenze SMN" in "".join(t.text or "" for t in tbl.iter(qn("w:t"))):
            tbl.getparent().remove(tbl)
    buf = io.BytesIO()
    doc.save(buf)
    fs = [f for f in check_rules(read_docx(buf.getvalue()), trip, deps.library) if f.rule == "transport_missing"]
    assert len(fs) == 1 and "Pisa Centrale" in fs[0].message and fs[0].day == 2


def test_unlinked_bookings_get_verify_marker_and_qc_rule(deps):
    from src.aigv2.models import Booking
    trip = day_trip().model_copy(update={"bookings": [
        Booking(name="Leaning Tower", day=M[1], time="11:00", source="voucher"),
        Booking(name="Uffizi Skip-the-line", day=M[2], time="10:00", source="voucher"),
        Booking(name="Cooking Class", day=date(2027, 4, 1), source="voucher")]})
    store = seeded(deps, trip)
    build_run(store, {}, deps)
    state = pipeline._load_plan(store)
    texts = {d.index: [m.render() for m in d.result.markers] for d in state.days}
    assert "Leaning Tower" in state.days[1].booking_by_item
    assert not any("Leaning Tower" in t for t in texts[2])
    assert any(t.startswith("⚠ VERIFY: booking 'Uffizi Skip-the-line' at 10:00") for t in texts[3])
    assert any("'Cooking Class'" in t and "2027-04-01" in t for t in texts[1])
    docx = store.load_bytes("guide.docx")
    assert "Uffizi Skip-the-line" in day_text(docx, 3)
    doc = Document(io.BytesIO(docx))
    for p in list(doc.element.body.iter(qn("w:p"))):
        if "Uffizi Skip-the-line" in "".join(t.text or "" for t in p.iter(qn("w:t"))):
            p.getparent().remove(p)
    buf = io.BytesIO()
    doc.save(buf)
    fs = [f for f in check_rules(read_docx(buf.getvalue()), trip, deps.library) if f.rule == "booking_missing"]
    assert [(f.day, "Uffizi Skip-the-line" in f.message) for f in fs] == [(3, True)]
