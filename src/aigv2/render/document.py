import io
import math
import re
from datetime import date, datetime, time, timedelta
from typing import Optional

from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from pydantic import BaseModel, Field

from src.aigv2.hours import format_for_day
from src.aigv2.library import norm_name
from src.aigv2.models import Booking, HotelStay, Marker, Place, TransportLeg, TripPlan
from src.aigv2.planning.important_places import ImportantPlaces
from src.aigv2.planning.meals import ChosenRestaurant, MealPick
from src.aigv2.planning.slots import DayResult, MealSlot, ScheduledItem
from src.aigv2.planning.travel import Leg, drive_day_legs, leg_markers, leg_options
from src.aigv2.render.blocks import (
    add_card, add_eyebrow, add_internal_link, add_link, add_marker, add_page_break, add_rule, add_slot_heading,
)
from src.aigv2.render.sections_render import (
    heading1, marker_lines, minutes_text, place_markers, render_important, render_section, render_thank_you,
    search_url,
)
from src.aigv2.render.theme import ASSETS, GOLD, PPR_ORDER, new_document, ordered_insert
from src.aigv2.writing.days import AttractionText, DayText
from src.aigv2.writing.sections import SectionContent

BEFORE_YOU_GO = [("packing", "Tailored Packing List"), ("money", "Money & Payments"),
                 ("connectivity", "Mobile Connectivity & SIM Guide"), ("health", "Health & Vaccination")]
WHILE_THERE = [("important", "Important Places Around Your Stays"), ("getting_around", "Getting Around"),
               ("dishes", "Must-Try Local Dishes"), ("souvenirs", "Souvenir Shopping Guide"),
               ("etiquette", "Cultural Etiquette & Local Phrases"), ("safety", "Safety & Emergency Contacts")]
_PORT = {"flight": "Airport", "train": "Station", "boat": "Harbor", "bus": "Bus station", "other": "Arrival point"}
_REACH = {"flight": "Leave for the airport by", "train": "Reach the station by", "boat": "Reach the harbor by"}
_SLOT_TITLE = {"morning": "Morning", "afternoon": "Afternoon", "evening": "Evening"}
_SUNSET = re.compile(r"sunset", re.I)
_BAD_FILE_CHARS = re.compile(r'[\\/:*?"<>|]')


class RenderDay(BaseModel):
    index: int
    date: date
    city: str
    text: DayText
    result: DayResult
    meals: list[MealPick] = Field(default_factory=list)
    legs: dict[str, Leg] = Field(default_factory=dict)
    places: dict[str, Optional[Place]] = Field(default_factory=dict)
    booking_by_item: dict[str, Booking] = Field(default_factory=dict)
    price_by_restaurant: dict[str, str | Marker] = Field(default_factory=dict)
    arrival: Optional[TransportLeg] = None
    departure: Optional[TransportLeg] = None
    prev_city: Optional[str] = None
    utc_offset_hours: Optional[float] = None
    library_attractions: dict[str, dict] = Field(default_factory=dict)


class GuideData(BaseModel):
    trip: TripPlan
    title: str
    days: list[RenderDay]
    sections: dict[str, Optional[SectionContent]]
    important: list[ImportantPlaces]
    trip_scope: str
    modes: list[str]
    domestic: bool


def day_header(index: int, d: date, title: str) -> str:
    return f"Day {index} – {d:%a}, {d.day} {d:%b} | {title}"


def destinations(trip: TripPlan) -> list[str]:
    return list(dict.fromkeys([h.city for h in trip.hotels] or [d.city for d in trip.days]))


def output_filename(trip: TripPlan) -> str:
    client = trip.client_names[0] if trip.client_names else "Client"
    name = f"All Inclusive Guide - {client} - {', '.join(destinations(trip))}"
    return _BAD_FILE_CHARS.sub("-", name) + ".docx"


def sunset_time(lat: float, lng: float, d: date, utc_offset_hours: float) -> Optional[time]:
    """NOAA approximate solar equations; None during polar day or night."""
    g = 2 * math.pi / 365 * (d.timetuple().tm_yday - 1)
    eq = 229.18 * (0.000075 + 0.001868 * math.cos(g) - 0.032077 * math.sin(g)
                   - 0.014615 * math.cos(2 * g) - 0.040849 * math.sin(2 * g))
    decl = (0.006918 - 0.399912 * math.cos(g) + 0.070257 * math.sin(g) - 0.006758 * math.cos(2 * g)
            + 0.000907 * math.sin(2 * g) - 0.002697 * math.cos(3 * g) + 0.00148 * math.sin(3 * g))
    la = math.radians(lat)
    c = math.cos(math.radians(90.833)) / (math.cos(la) * math.cos(decl)) - math.tan(la) * math.tan(decl)
    if not -1 <= c <= 1:
        return None
    minutes = round(720 - 4 * (lng - math.degrees(math.acos(c))) - eq + utc_offset_hours * 60) % 1440
    return time(minutes // 60, minutes % 60)


def _verify(text: str) -> Marker:
    return Marker(kind="VERIFY", text=text)


def _hm(dt: datetime) -> str:
    return dt.strftime("%H:%M")


class _DayRenderer:
    def __init__(self, doc, data: GuideData, rd: RenderDay, walk_max: int):
        self.doc, self.data, self.rd, self.walk_max = doc, data, rd, walk_max
        self.self_drive = "self_drive" in data.modes
        self.texts = {norm_name(a.name): a for a in rd.text.attractions}
        self.rtexts = {norm_name(r.name): r for r in rd.text.restaurants}
        self.plan = next((p for p in data.trip.days if p.date == rd.date), None)
        self.free_time_used = False

    def _stay(self, *, arriving: bool) -> Optional[HotelStay]:
        d, stays = self.rd.date, self.data.trip.hotels
        exact = [s for s in stays if (s.check_in if arriving else s.check_out) == d]
        cover = [s for s in stays if s.check_in <= d <= s.check_out]
        return (exact or cover or [None])[0]

    def _cruise(self) -> Optional[HotelStay]:
        d = self.rd.date
        return next((s for s in self.data.trip.hotels if s.kind == "cruise" and s.check_in < d < s.check_out), None)

    def _url(self, name: str) -> str:
        p = self.rd.places.get(name)
        return p.maps_url if p is not None and p.match_ok else search_url(name, self.rd.city)

    def _options(self, leg: Leg) -> list[str]:
        return [label if label.startswith("about") else f"{label} · {m} min"
                for label, m in leg_options(leg, self.data.modes, self.walk_max)]

    def render(self) -> None:
        rd, doc = self.rd, self.doc
        heading1(doc, day_header(rd.index, rd.date, rd.text.title), f"day_{rd.index}")
        doc.add_paragraph(rd.text.intro)
        for tip in rd.result.tips:
            doc.add_paragraph(tip, style="List Bullet")
        marker_lines(doc, rd.result.markers)
        cruise = self._cruise()
        arr, dep = rd.arrival, rd.departure
        early_dep = bool(dep and arr and (dep is arr or (dep.depart and arr.arrive and dep.depart <= arr.arrive)))
        if rd.prev_city and norm_name(rd.prev_city) != norm_name(rd.city) and not cruise:
            marker_lines(doc, [_verify("inter-city travel day logistics")])
        if cruise:
            self._port_day(arr, dep)
        else:
            if early_dep:
                self._departure(dep, transport_card=dep is not arr)
            if arr:
                self._arrival(arr)
        self._drive_stops()
        self._timeline()
        if rd.text.free_time_idea and not self.free_time_used:
            doc.add_paragraph(rd.text.free_time_idea)
        final = bool(dep and not early_dep and not cruise)
        if final:
            self._departure(dep, transport_card=True)
        else:
            self._overnight(cruise)

    def _overnight(self, cruise: Optional[HotelStay]) -> None:
        d = self.rd.date
        ship = cruise or next((s for s in self.data.trip.hotels
                               if s.kind == "cruise" and s.check_in <= d < s.check_out), None)
        text = f"Overnight on board {ship.hotel}" if ship else f"Overnight in {self.rd.city}"
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        run.italic = True

    def _port_day(self, arr: Optional[TransportLeg], dep: Optional[TransportLeg]) -> None:
        if arr and arr.arrive:
            self.doc.add_paragraph(f"Docks at {arr.to_place or self.rd.city} at {_hm(arr.arrive)}")
        if dep and dep.depart:
            self.doc.add_paragraph(f"Back on board by {_hm(dep.depart - timedelta(hours=1))}")
        if (arr and arr.arrive) or (dep and dep.depart):
            marker_lines(self.doc, [_verify("docking and all-aboard times")])

    def _leave_by(self, leg: TransportLeg) -> Optional[datetime]:
        if not leg.depart:
            return None
        if leg.kind == "flight":
            delta = timedelta(hours=2 if self.data.domestic else 3)
        elif leg.kind == "train":
            delta = timedelta(minutes=30)
        else:
            delta = timedelta(hours=1)
        return leg.depart - delta

    def _transfer_card(self, title: str, url: str, leg: Optional[Leg], body: Optional[str],
                       place: Optional[Place]) -> None:
        if leg is None:
            facts = [("TRANSFER", _verify("transfer time")), ("COST", _verify("transfer cost"))]
            markers = []
        else:
            opts = self._options(leg)
            facts = [("TRANSFER", opts or _verify("transfer time")), ("COST", _verify("transfer cost"))]
            markers = leg_markers(leg, self.data.modes, self.walk_max)
        add_card(self.doc, title, url, facts, body, [])
        marker_lines(self.doc, markers + place_markers(place))

    def _transport_card(self, t: TransportLeg, link_to: str) -> None:
        kind = {"flight": "Flight", "train": "Train", "boat": "Boat", "bus": "Bus"}.get(t.kind, "Transfer")
        title = f"{kind}: {t.from_place} to {t.to_place}" if t.from_place and t.to_place else kind
        facts: list = []
        if t.operator:
            facts.append(("OPERATOR", t.operator))
        if t.number:
            facts.append(("SERVICE", t.number))
        if t.seat:
            facts.append(("CLASS / SEAT", t.seat))
        if t.depart:
            facts.append(("DEPARTS", f"{_hm(t.depart)} · {t.from_place}" if t.from_place else _hm(t.depart)))
        if t.arrive:
            facts.append(("ARRIVES", f"{_hm(t.arrive)} · {t.to_place}" if t.to_place else _hm(t.arrive)))
        if t.kind == "train" and t.depart and t.arrive and t.arrive > t.depart:
            mins = int((t.arrive - t.depart).total_seconds() // 60)
            facts.append(("DURATION", [minutes_text(mins), _verify("duration (check time zones)")]))
        if t.kind == "train" and t.ref and t.platform:
            facts.append(("PLATFORM", t.platform))
        if t.ref:
            facts.append(("REFERENCE", t.ref))
        add_card(self.doc, title, search_url(link_to or t.to_place or t.from_place), facts, None, [])

    def _arrival(self, arr: TransportLeg) -> None:
        stay = self._stay(arriving=True)
        if arr is not self.rd.departure:
            self._transport_card(arr, arr.to_place)
        else:
            self._transport_card(arr, arr.from_place)
        if stay is None:
            marker_lines(self.doc, [_verify(f"hotel for {self.rd.date.isoformat()}")])
            return
        port = arr.to_place or _PORT[arr.kind]
        body = None
        check_in = stay.check_in_time or time(15)
        if arr.arrive and stay.check_in == self.rd.date and arr.arrive.time() < check_in:
            body = ("You arrive before check-in. If your room is not ready yet, ask the front desk to hold your "
                    "luggage while you head out.")
        self._transfer_card(f"{port} → {stay.hotel}", self._url(stay.hotel), self.rd.legs.get("arrival"), body,
                            self.rd.places.get(stay.hotel))

    def _departure(self, dep: TransportLeg, transport_card: bool) -> None:
        by = self._leave_by(dep)
        if by:
            self.doc.add_paragraph(f"{_REACH.get(dep.kind, 'Leave by')} {_hm(by)}")
        marker_lines(self.doc, [_verify("departure time")])
        stay = self._stay(arriving=False)
        port = dep.from_place or _PORT[dep.kind]
        if stay:
            self._transfer_card(f"{stay.hotel} → {port}", search_url(port, self.rd.city),
                                self.rd.legs.get("departure"), None, None)
        else:
            marker_lines(self.doc, [_verify(f"hotel to leave from on {self.rd.date.isoformat()}")])
        if transport_card:
            self._transport_card(dep, dep.from_place)

    def _drive_stops(self) -> None:
        if not self.plan or not self.plan.drive_stops:
            return
        _, minutes, km = drive_day_legs(self.rd.city, self.plan.drive_stops)
        add_eyebrow(self.doc, "Today's drive")
        for s in self.plan.drive_stops:
            self.doc.add_paragraph(f"{s.name} — {s.km_from_prev:g} km · {minutes_text(s.minutes_from_prev)}",
                                   style="List Bullet")
        self.doc.add_paragraph(f"Total today: {km:g} km · {minutes_text(minutes)}")

    def _timeline(self) -> None:
        entries = [(s.start, 0, i, s) for i, s in enumerate(self.rd.result.scheduled)]
        entries += [(m.start, 1, i, m) for i, m in enumerate(self.rd.result.meals)]
        slot = None
        for _, _, _, e in sorted(entries, key=lambda x: x[:3]):
            if isinstance(e, ScheduledItem):
                if e.slot != slot:
                    slot = e.slot
                    add_slot_heading(self.doc, _SLOT_TITLE[slot])
                self._item(e)
            else:
                self._meal(e)

    def _travel_line(self, name: str) -> None:
        leg = self.rd.legs.get(name)
        if leg is None:
            return
        opts = self._options(leg)
        if opts:
            self.doc.add_paragraph("Getting there: " + "  |  ".join(opts))
        if leg.transit_steps and any("public transport" in o for o in opts):
            st = leg.transit_steps
            self.doc.add_paragraph(f"Public transport: {' → '.join(s.line for s in st)}, alight {st[-1].alight}, "
                                   f"{sum(s.stops for s in st)} stops")
        marker_lines(self.doc, leg_markers(leg, self.data.modes, self.walk_max))

    def _parking(self, lib: dict):
        v = lib.get("parking")
        return v if isinstance(v, str) and v.strip() else _verify("parking")

    def _item(self, s: ScheduledItem) -> None:
        it, rd = s.item, self.rd
        if it.kind == "free_time":
            add_eyebrow(self.doc, "Free time")
            self.doc.add_paragraph(rd.text.free_time_idea or it.name)
            self.free_time_used = True
            return
        self._travel_line(it.name)
        at: Optional[AttractionText] = self.texts.get(norm_name(it.name))
        p = rd.places.get(it.name)
        found = p is not None and p.match_ok
        b = rd.booking_by_item.get(it.name)
        lib = rd.library_attractions.get(it.name, {})
        facts: list = []
        if it.pickup_time:
            facts.append(("PICKUP", it.pickup_time + (f" · {it.pickup_place}" if it.pickup_place else "")))
        if b:
            if not b.timed:
                facts.append(("YOUR BOOKING", "Valid all day"))
            elif b.time:
                facts.append(("YOUR BOOKING", b.time + (f"–{b.end_time}" if b.end_time else "")))
            else:
                facts.append(("YOUR BOOKING", _verify("booking time")))
            if b.meeting_point:
                facts.append(("MEETING POINT", b.meeting_point))
        elif it.time and not it.pickup_time:
            facts.append(("TIME", it.time))
        duration = (b.duration if b else "") or lib.get("recommended_duration") or lib.get("duration")
        if duration:
            facts.append(("DURATION", duration))
        elif at and at.duration:
            facts.append(("DURATION", [at.duration, _verify("duration")]))
        if found and it.kind != "transfer":
            facts.append(("OPENING HOURS", format_for_day(p.hours, rd.date) if p.hours else _verify("opening hours")))
        for label, value, what in (("ENTRY FEE", at and at.entry_fee, "entry fee"),
                                   ("DRESS CODE", at and at.dress_code, "dress code"),
                                   ("BEST TIME", at and at.best_time, "best time"),
                                   ("TIP", at and at.tips, "tip")):
            if value:
                facts.append((label, [value, _verify(what)]))
        if self.self_drive:
            facts.append(("PARKING", self._parking(lib)))
        sunset = self._sunset(it, at, p if found else None)
        if sunset:
            facts.append(("SUNSET", f"about {sunset:%H:%M}"))
        pills = (["Optional"] if it.optional else []) + (["Included in your package"] if b and b.included_in_package else [])
        url = p.maps_url if found else search_url(it.name, rd.city)
        table = add_card(self.doc, it.name, url, facts, at.description if at else None, pills)
        cell = table.cell(0, 0)
        if found and p.website and at and at.entry_fee:
            lp = cell.add_paragraph("Tickets and entry fees: ")
            add_link(lp, "official website", p.website)
        for instr in (b.instructions if b else []):
            cell.add_paragraph(instr, style="List Bullet")
        markers = [] if found else [_verify("place not found on Google")]
        if p is not None and p.status not in ("OPERATIONAL", "UNKNOWN"):
            markers.append(_verify(f"Google lists this place as {p.status.replace('_', ' ').lower()}"))
        marker_lines(self.doc, markers + place_markers(p))

    def _sunset(self, it, at: Optional[AttractionText], p: Optional[Place]) -> Optional[time]:
        if self.rd.utc_offset_hours is None:
            return None
        blob = " ".join(x for x in (it.name, it.note, at and at.description, at and at.tips) if x)
        if not _SUNSET.search(blob):
            return None
        where = p or self.rd.places.get(getattr(self._stay(arriving=True), "hotel", ""))
        if where is None:
            return None
        return sunset_time(where.lat, where.lng, self.rd.date, self.rd.utc_offset_hours)

    def _meal(self, slot: MealSlot) -> None:
        meal = slot.meal.capitalize()
        if slot.included:
            p = self.doc.add_paragraph()
            p.add_run(f"{meal} — Included in your package").bold = True
            return
        pick = next((m for m in self.rd.meals if m.slot.meal == slot.meal and m.slot.start == slot.start), None) \
            or next((m for m in self.rd.meals if m.slot.meal == slot.meal), None)
        if pick is None:
            add_slot_heading(self.doc, f"{meal} near {self.rd.city}")
            where = f"{meal}, Day {self.rd.index}, near {self.rd.city}"
            marker_lines(self.doc, [Marker(kind="ADD_RESTAURANT", text=where)] * 3)
            return
        if slot.unanchored or not slot.anchors:
            add_slot_heading(self.doc, f"{meal} near {self.rd.city}")
            for c in pick.restaurants:
                self._restaurant(c)
            marker_lines(self.doc, pick.markers)
            return
        rest, markers = list(pick.restaurants), list(pick.markers)
        for i, (name, _, count) in enumerate(slot.anchors):
            add_slot_heading(self.doc, f"{meal} near {name}")
            where = f"{meal}, Day {pick.day_index}, near {name}"
            mine = [m for m in markers if m.kind == "ADD_RESTAURANT" and m.text == where]
            take = len(rest) if i == len(slot.anchors) - 1 else max(count - len(mine), 0)
            for c in rest[:take]:
                self._restaurant(c)
            rest = rest[take:]
            marker_lines(self.doc, mine)
            markers = [m for m in markers if m not in mine]
        marker_lines(self.doc, markers)

    def _restaurant(self, c: ChosenRestaurant) -> None:
        lib, name = c.lib, c.lib.get("name", c.place.name)
        facts: list = []
        cuisine = lib.get("cuisine_type")
        if cuisine:
            facts.append(("CUISINE", ", ".join(cuisine) if isinstance(cuisine, list) else str(cuisine)))
        facts.append(("PRICE", self.rd.price_by_restaurant.get(name) or Marker(kind="PRICE_UNKNOWN", text="")))
        facts.append(("OPENING HOURS", [c.hours_text, _verify("opening hours from the library")]
                      if c.hours_verify else c.hours_text))
        facts.append(("GETTING THERE", f"{c.travel_label} · {c.travel_min} min"))
        dishes = lib.get("must_try_dishes")
        if isinstance(dishes, list) and dishes:
            facts.append(("MUST TRY", ", ".join(str(x) for x in dishes[:3])))
        rt = self.rtexts.get(norm_name(name))
        if rt and rt.best_for:
            facts.append(("BEST FOR", rt.best_for))
        if self.self_drive:
            facts.append(("PARKING", self._parking(lib)))
        pills = ["Vegetarian friendly"] if lib.get("vegetarian_friendly") is True else []
        add_card(self.doc, name, c.place.maps_url or search_url(name, self.rd.city), facts,
                 rt.ambience if rt else None, pills)
        marker_lines(self.doc, place_markers(c.place))


def _present(data: GuideData, key: str) -> bool:
    if key == "important":
        return bool(data.important)
    s = data.sections.get(key)
    return s is not None and bool(s.items or s.by_country)


def _cover(doc) -> None:
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(0)
    sec.header_distance = sec.footer_distance = Cm(0)
    sec.different_first_page_header_footer = True
    p = doc.add_paragraph()
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(ASSETS / "cover_a4.png"), width=Cm(21.0), height=Cm(29.7))
    new = doc.add_section(WD_SECTION.NEW_PAGE)
    holder = doc.paragraphs[-1]._p
    sect = holder.pPr.sectPr
    ordered_insert(p._p.get_or_add_pPr(), sect, PPR_ORDER)
    holder.getparent().remove(holder)
    new.left_margin = new.right_margin = new.top_margin = new.bottom_margin = Cm(2.0)
    new.header_distance = new.footer_distance = Cm(1.0)
    new.different_first_page_header_footer = False


def _title_page(doc, data: GuideData, dest: str, toc: list[tuple[str, list[tuple[str, str]]]]) -> None:
    lp = doc.add_paragraph()
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lp.add_run().add_picture(str(ASSETS / "logo.png"), width=Cm(6))
    kicker = add_eyebrow(doc, "Your All-Inclusive Guide")
    for run in kicker.runs:
        run.font.color.rgb = RGBColor.from_string(GOLD)
        run.font.italic = True
    doc.add_paragraph(data.title, style="Title")
    doc.add_paragraph(f"All Inclusive Guide – {dest}")
    add_rule(doc)
    for group, entries in toc:
        if not entries:
            continue
        add_eyebrow(doc, group)
        for text, bookmark in entries:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            add_internal_link(p, text, bookmark)


def render_guide(data: GuideData, settings) -> bytes:
    doc = new_document()
    dest = ", ".join(destinations(data.trip))
    _cover(doc)
    first = data.trip.client_names[0].split()[0] if data.trip.client_names else ""
    doc.add_paragraph(f"Dear {first},")
    p = doc.add_paragraph()
    add_marker(p, Marker(kind="WRITE_LETTER", text=""))
    add_page_break(doc)

    before = [(k, t) for k, t in BEFORE_YOU_GO if _present(data, k)]
    while_ = [(k, t) for k, t in WHILE_THERE if _present(data, k)]
    toc = [("DAY BY DAY", [(day_header(d.index, d.date, d.text.title), f"day_{d.index}") for d in data.days]),
           ("BEFORE YOU GO", [(t, f"sec_{k}") for k, t in before]),
           ("WHILE YOU'RE THERE", [(t, f"sec_{k}") for k, t in while_])]
    _title_page(doc, data, dest, toc)

    for rd in data.days:
        _DayRenderer(doc, data, rd, settings.walk_max_min).render()
    for key, title in before + while_:
        if key == "important":
            render_important(doc, title, f"sec_{key}", data.important)
        else:
            render_section(doc, title, f"sec_{key}", data.sections[key])
    render_thank_you(doc, dest)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
