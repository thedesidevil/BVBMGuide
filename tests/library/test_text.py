import pytest
from docx import Document

from src.library.ingest_core.text import UnreadableFile, read_text, text_hash


def test_read_text_includes_table_rows_in_body_order(tmp_path):
    doc = Document()
    doc.add_paragraph("Before table")
    table = doc.add_table(rows=2, cols=2)
    for r, row in enumerate([["Day 1", "Arrive"], ["Day 2", "Museum"]]):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    doc.add_paragraph("After table")
    p = tmp_path / "a.docx"
    doc.save(p)
    lines = [l for l in read_text(p).splitlines() if l.strip()]
    assert lines == ["Before table", "Day 1 | Arrive", "Day 2 | Museum", "After table"]


def test_read_text_unreadable_raises(tmp_path):
    p = tmp_path / "bad.docx"
    p.write_text("not a docx")
    with pytest.raises(UnreadableFile):
        read_text(p)


def test_read_text_unsupported_suffix_raises(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("x")
    with pytest.raises(UnreadableFile):
        read_text(p)


def test_text_hash_ignores_whitespace():
    assert text_hash("a  b\n c") == text_hash("a b c")
    assert text_hash("a b") != text_hash("a c")
