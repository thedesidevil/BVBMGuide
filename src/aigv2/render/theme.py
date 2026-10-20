from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ASSETS = Path(__file__).resolve().parent.parent / "assets"

NAVY, ROYAL, GOLD, SLATE = "0C1F3E", "2F5FA8", "C68A32", "41536E"
POWDER, FROST, CARD_BORDER, CARD_FILL = "AEC7E7", "EEF4FB", "DCE6F2", "FBFDFF"
HEADING_FONT, BODY_FONT = "Cormorant Garamond", "Georgia"
TAGLINE = "Your Vacation · Your Story · Your Way"


def set_font(style_or_run, name: str) -> None:
    style_or_run.font.name = name
    rpr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run, "element") else style_or_run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn(f"w:{attr}"), name)
    for attr in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
        if rfonts.get(qn(f"w:{attr}")) is not None:
            del rfonts.attrib[qn(f"w:{attr}")]


def _style(styles, name, font, size, color, bold=False, italic=False, before=0, after=6):
    s = styles[name]
    set_font(s, font)
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor.from_string(color)
    s.font.bold = bold
    s.font.italic = italic
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    return s


def _field(paragraph, instr: str):
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        run = paragraph.add_run()
        _style_footer_run(run)
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


def _tab(kind, pos):
    t = OxmlElement("w:tab")
    t.set(qn("w:val"), kind)
    t.set(qn("w:pos"), str(pos))
    return t


def _style_footer_run(run):
    set_font(run, HEADING_FONT)
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = RGBColor.from_string(SLATE)


def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.header_distance = sec.footer_distance = Cm(1.0)

    st = doc.styles
    _style(st, "Normal", BODY_FONT, 10.5, SLATE, after=6)
    _style(st, "Title", HEADING_FONT, 34, NAVY, after=12)
    _style(st, "Heading 1", HEADING_FONT, 24, NAVY, before=18, after=8)
    _style(st, "Heading 2", HEADING_FONT, 15, GOLD, italic=True, before=14, after=6)
    _style(st, "Heading 3", HEADING_FONT, 15, NAVY, bold=True, before=0, after=4)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        ppr = st[name].element.get_or_add_pPr()
        for bdr in ppr.findall(qn("w:pBdr")):
            ppr.remove(bdr)

    hp = sec.header.paragraphs[0]
    hp.add_run().add_picture(str(ASSETS / "logo.png"), width=Cm(3))

    fp = sec.footer.paragraphs[0]
    width = sec.page_width - sec.left_margin - sec.right_margin
    fp.style = doc.styles["Footer"]
    ppr = fp._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    for pos in (4680, 9360):
        tabs.append(_tab("clear", pos))
    tabs.append(_tab("right", int(width / 635)))
    ppr.append(tabs)
    _style_footer_run(fp.add_run(TAGLINE))
    _style_footer_run(fp.add_run("\t"))
    _field(fp, "PAGE")
    return doc
