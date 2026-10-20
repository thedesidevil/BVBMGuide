from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from pydantic import BaseModel, Field

FIELDS: list[tuple[str, bool]] = [
    ("Client name(s)", True),
    ("Number of travellers (adults)", True),
    ("Kids' ages", True),
    ("Diet (restrictions, allergies, likes)", True),
    ("Travel mode(s)", True),
    ("Itinerary", True),
    ("Trip dates", False),
    ("Cities and nights", False),
    ("Hotel per city", False),
    ("Flights / trains", False),
    ("Confirmed bookings (tours, shows, times)", False),
    ("Seniors / mobility notes", False),
    ("Budget level", False),
    ("Occasion", False),
    ("Must-dos / avoid", False),
    ("Additional notes", False),
]

SEE_ATTACHMENTS = {"see attachments", "see attached", "use information from the attached files"}
_FONT = "Georgia"


class IntakeFields(BaseModel):
    values: dict[str, str] = Field(default_factory=dict)
    see_attachments: set[str] = Field(default_factory=set)
    missing_required: list[str] = Field(default_factory=list)
    is_template: bool = False


def _norm(label: str) -> str:
    return label.replace("*", "").strip().lower()


def _borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "999999")
        borders.append(el)
    tbl_pr.append(borders)


def _set(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.name = _FONT
    run.font.size = Pt(10.5)


def build_template(path: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = _FONT
    doc.add_heading("AIG Intake", level=1)
    doc.add_paragraph(
        "Fill in the answers below. Write 'see attachments' for anything in the attached files. "
        "Fields marked * are required."
    )
    table = doc.add_table(rows=1, cols=2)
    _borders(table)
    _set(table.rows[0].cells[0], "Field", bold=True)
    _set(table.rows[0].cells[1], "Your answer", bold=True)
    for label, required in FIELDS:
        row = table.add_row()
        _set(row.cells[0], label + (" *" if required else ""), bold=True)
        _set(row.cells[1], "")
    for row in table.rows:
        row.cells[0].width = Pt(190)
        row.cells[1].width = Pt(300)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def parse_template(path: Path) -> IntakeFields:
    doc = Document(str(path))
    known = {_norm(label): label for label, _ in FIELDS}
    required = [label for label, req in FIELDS if req]
    for table in doc.tables:
        found: dict[str, str] = {}
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            label = known.get(_norm(row.cells[0].text))
            if label:
                found[label] = row.cells[1].text.strip()
        if not found:
            continue
        see = {k for k, v in found.items() if v.lower() in SEE_ATTACHMENTS}
        values = {k: v for k, v in found.items() if v and k not in see}
        missing = [r for r in required if r not in values and r not in see]
        return IntakeFields(values=values, see_attachments=see, missing_required=missing, is_template=True)
    return IntakeFields()
